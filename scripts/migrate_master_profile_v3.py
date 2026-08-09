#!/usr/bin/env python3
"""Migration sûre du master profile vers le schéma CV V3.

Usage :
    python scripts/migrate_master_profile_v3.py data/reference/master_profile.xlsx

La migration est idempotente. Elle sauvegarde le classeur avant toute écriture,
écrit via un fichier temporaire, et conserve l'ancienne feuille sous le nom
``leadership_legacy``.
"""

from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass
from datetime import datetime
import hashlib
import os
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata

from openpyxl import load_workbook


SAP_EASY_ACCESS = "SAP Easy Access"
SAP_EWM = "SAP S/4HANA – Extended Warehouse Management (EWM)"


@dataclass(frozen=True)
class MigrationResult:
    workbook: str
    migrated_experiences: int
    normalized_sap_entries: int
    legacy_sheet: str | None
    backup: str | None
    changed: bool

    def to_dict(self) -> dict:
        return asdict(self)


def _headers(sheet) -> dict[str, int]:
    return {
        str(cell.value).strip().casefold(): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value not in (None, "")
    }


def _ensure_headers(sheet, names: list[str]) -> dict[str, int]:
    headers = _headers(sheet)
    for name in names:
        if name.casefold() not in headers:
            sheet.cell(row=1, column=sheet.max_column + 1, value=name)
            headers[name.casefold()] = sheet.max_column
    return headers


def _value(sheet, row: int, headers: dict[str, int], *names: str):
    for name in names:
        column = headers.get(name.casefold())
        if column:
            value = sheet.cell(row=row, column=column).value
            if value not in (None, ""):
                return value
    return ""


def _slug(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or "").casefold())
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "_", normalized).strip("_")[:48] or "experience"


def _legacy_key(company, role, dates, location) -> str:
    raw = "|".join(str(value or "").strip().casefold() for value in (company, role, dates, location))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]


def _stable_id(company, role, dates, location) -> str:
    key = _legacy_key(company, role, dates, location)
    return f"freelance_{_slug(f'{company}_{role}')}_{key[:8]}"


def _year_bounds(value) -> tuple[str, str]:
    years = re.findall(r"(?:19|20)\d{2}", str(value or ""))
    if not years:
        return "", ""
    return years[0], years[-1]


def _normalized(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _normalize_sap_certifications(workbook) -> int:
    if "certifications" not in workbook.sheetnames:
        return 0
    sheet = workbook["certifications"]
    headers = _headers(sheet)
    column = next(
        (headers[name] for name in ("cert_name", "certification_name", "certification", "name", "title") if name in headers),
        None,
    )
    if not column:
        return 0
    changed = 0
    easy_access_present = False
    ewm_present = False
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=column)
        name = _normalized(cell.value)
        replacement = None
        if name in {"sap easy access", "sap supply chain management"}:
            replacement = SAP_EASY_ACCESS
            easy_access_present = True
        elif name in {
            "sap 4hana ewm",
            "sap s 4hana ewm",
            "sap s 4hana extended warehouse management ewm",
            "sap supply chain management ewm 4hana",
        }:
            replacement = SAP_EWM
            ewm_present = True
        elif name == _normalized(SAP_EASY_ACCESS):
            easy_access_present = True
        elif name == _normalized(SAP_EWM):
            ewm_present = True
        if replacement and cell.value != replacement:
            cell.value = replacement
            changed += 1

    # Les deux formations demandées doivent exister, sans effacer les autres.
    for missing_name in (
        [] if easy_access_present else [SAP_EASY_ACCESS]
    ) + ([] if ewm_present else [SAP_EWM]):
        values = {header: "" for header in headers}
        values[next(name for name, index in headers.items() if index == column)] = missing_name
        if "issuer" in headers:
            values["issuer"] = "SAP"
        if "status" in headers:
            values["status"] = "verified"
        sheet.append([values.get(str(cell.value).strip().casefold(), "") for cell in sheet[1]])
        changed += 1
    return changed


def migrate_workbook(path: str | Path, *, backup_dir: str | Path | None = None) -> MigrationResult:
    workbook_path = Path(path).expanduser().resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)
    workbook = load_workbook(workbook_path)
    if "experiences" not in workbook.sheetnames:
        raise ValueError("La feuille experiences est obligatoire.")

    experiences = workbook["experiences"]
    experience_headers = _ensure_headers(
        experiences,
        [
            "experience_id",
            "company",
            "location",
            "job_title",
            "date_start",
            "date_end",
            "dates",
            *(f"truth_bullet_{index}" for index in range(1, 11)),
            "validated_memory",
            "tools_verified",
            "is_freelance",
            "source_legacy",
            "legacy_source_key",
        ],
    )

    existing_ids = {
        str(experiences.cell(row=row, column=experience_headers["experience_id"]).value or "").strip()
        for row in range(2, experiences.max_row + 1)
    }
    existing_keys = {
        str(experiences.cell(row=row, column=experience_headers["legacy_source_key"]).value or "").strip()
        for row in range(2, experiences.max_row + 1)
    }

    source_name = "leadership" if "leadership" in workbook.sheetnames else (
        "leadership_legacy" if "leadership_legacy" in workbook.sheetnames else None
    )
    migrated = 0
    if source_name:
        source = workbook[source_name]
        source_headers = _headers(source)
        for row in range(2, source.max_row + 1):
            company = _value(source, row, source_headers, "organisation", "organization", "company", "org", "project")
            role = _value(source, row, source_headers, "role", "position_title", "job_title", "title")
            dates = _value(source, row, source_headers, "dates", "year", "date", "date_range")
            location = _value(source, row, source_headers, "location", "city", "city_state", "lieu")
            bullets = [
                _value(source, row, source_headers, f"truth_bullet_{index}", f"bullet_{index}")
                for index in range(1, 11)
            ]
            if not any((company, role, dates, location, *bullets)):
                continue
            key = _legacy_key(company, role, dates, location)
            experience_id = _stable_id(company, role, dates, location)
            if key in existing_keys or experience_id in existing_ids:
                continue
            start, end = _year_bounds(dates)
            values = {
                "experience_id": experience_id,
                "company": company,
                "location": location,
                "job_title": role,
                "date_start": start,
                "date_end": end,
                "dates": dates,
                "validated_memory": _value(source, row, source_headers, "validated_memory"),
                "tools_verified": _value(source, row, source_headers, "tools_verified", "tools"),
                "is_freelance": True,
                "source_legacy": "leadership",
                "legacy_source_key": key,
            }
            for index, bullet in enumerate(bullets, start=1):
                values[f"truth_bullet_{index}"] = bullet
            # Copie aussi les colonnes homonymes présentes dans les deux schémas.
            for name in experience_headers:
                if name not in values and name in source_headers:
                    values[name] = _value(source, row, source_headers, name)
            experiences.append(
                [values.get(str(cell.value).strip().casefold(), "") for cell in experiences[1]]
            )
            migrated += 1
            existing_ids.add(experience_id)
            existing_keys.add(key)

    renamed = None
    if "leadership" in workbook.sheetnames and "leadership_legacy" not in workbook.sheetnames:
        workbook["leadership"].title = "leadership_legacy"
        renamed = "leadership_legacy"

    sap_updates = _normalize_sap_certifications(workbook)
    changed = bool(migrated or renamed or sap_updates)
    backup_path = None
    if changed:
        backup_root = Path(backup_dir) if backup_dir else workbook_path.parent / "backups"
        backup_root.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        backup_path = backup_root / f"{workbook_path.stem}_before_v3_{timestamp}.xlsx"
        shutil.copy2(workbook_path, backup_path)
        with tempfile.NamedTemporaryFile(dir=workbook_path.parent, suffix=".xlsx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        try:
            workbook.save(temporary_path)
            os.replace(temporary_path, workbook_path)
        finally:
            temporary_path.unlink(missing_ok=True)

    return MigrationResult(
        workbook=str(workbook_path),
        migrated_experiences=migrated,
        normalized_sap_entries=sap_updates,
        legacy_sheet=renamed or ("leadership_legacy" if "leadership_legacy" in workbook.sheetnames else None),
        backup=str(backup_path) if backup_path else None,
        changed=changed,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbooks", nargs="+", help="Classeur(s) master_profile.xlsx à migrer")
    parser.add_argument("--backup-dir", help="Dossier de sauvegarde optionnel")
    arguments = parser.parse_args()
    for workbook in arguments.workbooks:
        print(migrate_workbook(workbook, backup_dir=arguments.backup_dir).to_dict())


if __name__ == "__main__":
    main()
