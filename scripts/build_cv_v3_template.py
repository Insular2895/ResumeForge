#!/usr/bin/env python3
"""Construit le template de production propre à partir du DOCX V3 fourni."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


FONT = "Calibri"


def _stabilize_floating_images(document) -> None:
    """Empêche une photo flottante de déplacer ou masquer le texte du CV.

    Le document V3 source utilise un ``wrapSquare`` rattaché au paragraphe du
    titre. Selon le contenu injecté, Quick Look et certains moteurs DOCX ne
    calculent pas tous cet habillage de la même façon. La photo reste flottante
    au même emplacement et visible au premier plan, mais ne participe plus au
    flux typographique.
    """
    for anchor in document._element.findall(".//" + qn("wp:anchor")):
        anchor.set("behindDoc", "0")
        wrap_elements = [
            child
            for child in list(anchor)
            if child.tag
            in {
                qn("wp:wrapNone"),
                qn("wp:wrapSquare"),
                qn("wp:wrapTight"),
                qn("wp:wrapThrough"),
                qn("wp:wrapTopAndBottom"),
            }
        ]
        if wrap_elements:
            first_index = anchor.index(wrap_elements[0])
            for element in wrap_elements:
                anchor.remove(element)
            anchor.insert(first_index, OxmlElement("wp:wrapNone"))


def _set_font(run, size: float = 10, bold: bool | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)


def _set_text(paragraph, text: str, *, size=10, bold=False, align=None) -> None:
    for run in paragraph.runs:
        if not run._element.findall(".//" + qn("w:drawing")):
            run.text = ""
    run = next(
        (candidate for candidate in paragraph.runs if not candidate._element.findall(".//" + qn("w:drawing"))),
        None,
    ) or paragraph.add_run()
    run.text = text
    _set_font(run, size=size, bold=bold)
    if align is not None:
        paragraph.alignment = align


def _reset_indentation(paragraph) -> None:
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._p
    element.getparent().remove(element)


def _canonical_text(text: str) -> str:
    compact = re.sub(r"\s+", " ", str(text or "")).strip()
    compact = compact.replace("COMPAGNY", "COMPANY")
    if "[[CV_HEADLINE]]" in compact:
        return "[[CV_HEADLINE]]"
    for index in (1, 2, 3):
        for field in ("COMPANY", "POSITION_TITLE", "BULLETS"):
            placeholder = f"[[EXP_{index}_{field}]]"
            if placeholder in compact:
                return placeholder
        if f"[[EXP_{index}_LOCATION]]" in compact or f"[[EXP_{index}_DATES]]" in compact:
            return f"[[EXP_{index}_LOCATION]] - [[EXP_{index}_DATES]]"
    if "[[CERTIFICATION_ENTRIES]]" in compact:
        return "[[CERTIFICATION_ENTRIES]]"
    if "[[TECHNICAL_SKILLS]]" in compact:
        return "Compétences techniques : [[TECHNICAL_SKILLS]]"
    return compact.rstrip(" -–—")


def build_template(source: str | Path, destinations: list[str | Path]) -> list[Path]:
    source_path = Path(source)
    if not source_path.exists():
        raise FileNotFoundError(source_path)
    document = Document(source_path)

    _stabilize_floating_images(document)

    body_section = document._element.body.sectPr
    for section_properties in list(document._element.findall(".//" + qn("w:pPr") + "/" + qn("w:sectPr"))):
        if section_properties is not body_section:
            section_properties.getparent().remove(section_properties)
    for break_element in list(document._element.findall(".//" + qn("w:br"))):
        break_element.getparent().remove(break_element)

    for paragraph in list(document.paragraphs):
        canonical = _canonical_text(paragraph.text)
        if not canonical or "[[Title_Job" in canonical:
            _remove_paragraph(paragraph)
            continue
        if not re.fullmatch(r"\[\[EXP_[123]_BULLETS\]\]", canonical):
            # Le fichier source contient plusieurs retraits directs invisibles
            # (dont 7,6 cm sur « Éducation »). Ils survivent à un simple
            # changement d'alignement et décentrent le rendu final.
            _reset_indentation(paragraph)
        normalized = canonical.casefold()
        if canonical == "[[CV_HEADLINE]]":
            _set_text(paragraph, canonical, size=12.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.paragraph_format.space_before = Pt(0)
        elif normalized.startswith("lucas pertusa"):
            # Conserve le lien e-mail et les runs mixtes du document source.
            if paragraph.runs:
                paragraph.runs[0].text = paragraph.runs[0].text.lstrip()
            for run in paragraph.runs:
                _set_font(run, size=11.5, bold=run.bold)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            # Laisse la photo ancrée se terminer avant le bloc Éducation afin
            # de conserver une respiration visuelle stable entre les blocs.
            paragraph.paragraph_format.space_after = Pt(40)
        elif normalized in {
            "éducation",
            "education",
            "formations & certifications",
            "expériences",
            "experiences",
            "compétences & langues",
            "skills & languages",
        }:
            _set_text(paragraph, canonical, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(2)
        elif re.fullmatch(r"\[\[EXP_[123]_COMPANY\]\]", canonical):
            _set_text(paragraph, canonical, size=10, bold=True)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(0)
        elif re.fullmatch(r"\[\[EXP_[123]_POSITION_TITLE\]\]", canonical):
            _set_text(paragraph, canonical, size=10, bold=True)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
        elif re.fullmatch(r"\[\[EXP_[123]_LOCATION\]\] - \[\[EXP_[123]_DATES\]\]", canonical):
            _set_text(paragraph, canonical, size=10, bold=False)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
        elif re.fullmatch(r"\[\[EXP_[123]_BULLETS\]\]", canonical):
            _set_text(paragraph, canonical, size=10, bold=False)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
        elif canonical == "[[CERTIFICATION_ENTRIES]]":
            _set_text(paragraph, canonical, size=10, bold=False)
            paragraph.paragraph_format.space_after = Pt(2)
        elif normalized.startswith("compétences techniques :") or normalized.startswith("langues :"):
            _set_text(paragraph, canonical, size=10, bold=False)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)
        else:
            _set_text(paragraph, canonical, size=10, bold=bool(any(run.bold for run in paragraph.runs)))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)

    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = FONT
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_font(run, size=10, bold=run.bold)

    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(10)
    section.left_margin = Mm(12.5)
    section.right_margin = Mm(12.5)
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)
    columns = body_section.find(qn("w:cols")) if body_section is not None else None
    if columns is None and body_section is not None:
        columns = OxmlElement("w:cols")
        body_section.append(columns)
    if columns is not None:
        columns.set(qn("w:num"), "1")
        for column in list(columns.findall(qn("w:col"))):
            columns.remove(column)

    # Certains éditeurs enregistrent des twips sous forme de flottants
    # (ex. ``358.999999``), alors que WordprocessingML exige un entier.
    # python-docx lève sinon une ValueError lors de la copie du format.
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for element in document._element.iter():
        for attribute, value in list(element.attrib.items()):
            if attribute.startswith(word_namespace) and re.fullmatch(r"-?\d+\.\d+", str(value)):
                element.set(attribute, str(round(float(value))))

    written = []
    for destination in destinations:
        path = Path(destination)
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)
        written.append(path)
    return written


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("destinations", nargs="+")
    args = parser.parse_args()
    for path in build_template(args.source, args.destinations):
        print(path)


if __name__ == "__main__":
    main()
