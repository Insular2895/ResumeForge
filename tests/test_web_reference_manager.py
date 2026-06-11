from pathlib import Path

import pytest
from docx import Document

from src.web import reference_manager


def _docx(path: Path, text: str) -> Path:
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def test_reference_status_reports_missing_and_ready(tmp_path, monkeypatch):
    specs = {
        "reference_cv": reference_manager.ReferenceSpec(
            key="reference_cv",
            label="CV de référence",
            destination=tmp_path / "reference_cv.txt",
            extensions=(".docx", ".md", ".txt"),
        )
    }
    monkeypatch.setattr(reference_manager, "REFERENCE_SPECS", specs)

    assert reference_manager.get_reference_statuses()["reference_cv"]["ready"] is False

    specs["reference_cv"].destination.write_text("CV", encoding="utf-8")

    assert reference_manager.get_reference_statuses()["reference_cv"]["ready"] is True


def test_reference_status_reports_invalid_existing_file(tmp_path, monkeypatch):
    invalid = tmp_path / "base_cv.docx"
    invalid.write_text("pas un document Word", encoding="utf-8")
    specs = {
        "cv_template": reference_manager.ReferenceSpec(
            key="cv_template",
            label="Template CV Word",
            destination=invalid,
            extensions=(".docx",),
            required_placeholders=("[[EXP_1_BULLETS]]",),
        )
    }
    monkeypatch.setattr(reference_manager, "REFERENCE_SPECS", specs)

    status = reference_manager.get_reference_statuses()["cv_template"]

    assert status["ready"] is False
    assert status["error"]


def test_replace_reference_rejects_wrong_extension(tmp_path, monkeypatch):
    spec = reference_manager.ReferenceSpec(
        key="master_profile",
        label="Profil Excel maître",
        destination=tmp_path / "master_profile.xlsx",
        extensions=(".xlsx",),
    )
    monkeypatch.setattr(reference_manager, "REFERENCE_SPECS", {"master_profile": spec})
    uploaded = tmp_path / "profile.txt"
    uploaded.write_text("bad", encoding="utf-8")

    with pytest.raises(ValueError, match="Extension"):
        reference_manager.replace_reference("master_profile", uploaded)


def test_replace_docx_template_requires_expected_placeholder(tmp_path, monkeypatch):
    destination = tmp_path / "base_cv.docx"
    spec = reference_manager.ReferenceSpec(
        key="cv_template",
        label="Template CV Word",
        destination=destination,
        extensions=(".docx",),
        required_placeholders=("[[EXP_1_BULLETS]]",),
    )
    monkeypatch.setattr(reference_manager, "REFERENCE_SPECS", {"cv_template": spec})
    uploaded = _docx(tmp_path / "uploaded.docx", "Aucun placeholder")

    with pytest.raises(ValueError, match="placeholder"):
        reference_manager.replace_reference("cv_template", uploaded)

    assert not destination.exists()


def test_replace_reference_keeps_previous_file_when_validation_fails(tmp_path, monkeypatch):
    destination = _docx(tmp_path / "base_letter.docx", "[[LM_FINAL_LETTER]]")
    original = destination.read_bytes()
    spec = reference_manager.ReferenceSpec(
        key="lm_template",
        label="Template LM Word",
        destination=destination,
        extensions=(".docx",),
        required_placeholders=("[[LM_FINAL_LETTER]]",),
    )
    monkeypatch.setattr(reference_manager, "REFERENCE_SPECS", {"lm_template": spec})
    invalid = _docx(tmp_path / "invalid.docx", "Invalide")

    with pytest.raises(ValueError):
        reference_manager.replace_reference("lm_template", invalid)

    assert destination.read_bytes() == original
