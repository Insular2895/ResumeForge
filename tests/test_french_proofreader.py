from docx import Document
import pytest

from src.letter.french_proofreader import (
    apply_known_french_corrections,
    enforce_french_docx,
    check_french_text,
)


def test_known_cv_typos_are_corrected_before_rendering():
    content = {
        "role": "**Devellopper web**",
        "certifications": ["Portflio and risk management"],
    }

    corrected = apply_known_french_corrections(content)

    assert corrected["role"] == "Développeur web"
    assert corrected["certifications"] == ["Portfolio and risk management"]


def test_french_checker_rejects_clear_spelling_error():
    report = check_french_text(
        "Cette lettre ne doit contenir aucune faute d'ortographe.",
    )

    assert report["status"] == "failed"
    assert any(issue["suggestion"] == "orthographe" for issue in report["issues"])


def test_french_checker_allows_professional_terms_and_proper_names():
    report = check_french_text(
        "Cuisinella recherche un profil ADV capable de produire un reporting fiable.",
    )

    assert report["status"] == "success"


def test_french_checker_allows_valid_logistics_plural():
    report = check_french_text(
        "Le suivi précis des acheminements sécurise les approvisionnements.",
    )

    assert report["status"] == "success"


def test_french_checker_allows_ellipsis_but_rejects_repeated_exclamation():
    assert check_french_text("Organisation, suivi, coordination...")["status"] == "success"
    assert check_french_text("Organisation impeccable!!!")["status"] == "failed"


def test_final_docx_auto_corrects_safe_known_errors(tmp_path):
    output_path = tmp_path / "CV.docx"
    document = Document()
    document.add_paragraph("Une **motivassion** solide , sans erreur.")
    document.save(output_path)

    enforce_french_docx(output_path, artifact_label="CV")

    rendered = Document(output_path)
    assert rendered.paragraphs[0].text == "Une motivation solide, sans erreur."


def test_final_docx_checks_headers_too(tmp_path):
    output_path = tmp_path / "CV.docx"
    document = Document()
    document.add_paragraph("Expérience professionnelle.")
    document.sections[0].header.paragraphs[0].text = "Organisation impeccable!!!"
    document.save(output_path)

    with pytest.raises(RuntimeError, match="ponctuation répétée"):
        enforce_french_docx(output_path, artifact_label="CV")

    assert not output_path.exists()
