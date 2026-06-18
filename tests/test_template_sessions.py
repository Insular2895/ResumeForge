import json
import zipfile
from pathlib import Path

from docx import Document

from src.web import template_sessions


def _docx(path: Path, paragraphs: list[str]) -> Path:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)
    return path


def test_create_template_session_extracts_json_and_generates_docx(tmp_path):
    pack = tmp_path / "pack"
    pack.mkdir()
    _docx(
        pack / "CV - Test.docx",
        [
            "Lucas Pertusa",
            "Acheteur & Coordinateur Supply Chain",
            "Paris | lucas@example.com",
            "Profil achats et supply.",
            "Expérience",
            "Adventis",
            "Acheteur junior",
            "2024 - 2025",
            "Sourcing et qualification fournisseurs",
            "Suivi lead time et OTIF",
            "Compétences",
            "Sourcing",
            "SRM",
        ],
    )
    _docx(
        pack / "LM - Test.docx",
        [
            "Objet : Candidature Acheteur",
            "Madame, Monsieur,",
            "Intro personnalisée.",
            "Argument achats.",
            "Argument supply.",
            "Je vous prie d'agréer.",
            "Lucas Pertusa",
        ],
    )

    session = template_sessions.create_template_session(pack, tmp_path / "sessions")

    root = tmp_path / "sessions" / session["session_id"]
    assert session["documents"]["cv"]["docx_filename"] == "CV_Lucas_Pertusa.docx"
    assert session["documents"]["lm"]["docx_filename"] == "Lettre_Motivation_Lucas_Pertusa.docx"
    assert json.loads((root / "cv_data.json").read_text(encoding="utf-8"))["name"] == "Lucas Pertusa"
    assert json.loads((root / "lm_data.json").read_text(encoding="utf-8"))["subject"] == "Objet : Candidature Acheteur"
    assert (root / "CV_Lucas_Pertusa.docx").exists()
    assert (root / "Lettre_Motivation_Lucas_Pertusa.docx").exists()


def test_update_session_persists_json_regenerates_docx_and_zip(tmp_path):
    session_root = tmp_path / "sessions"
    session = template_sessions.create_template_session_from_data(
        {
            "name": "Lucas Pertusa",
            "title": "Acheteur",
            "contact": "Paris",
            "profile": "Ancien profil",
            "skills": ["Sourcing"],
            "experiences": [
                {
                    "company": "Adventis",
                    "role": "Acheteur junior",
                    "dates": "2024 - 2025",
                    "bullets": ["Ancien bullet"],
                }
            ],
            "education": ["Bachelor"],
        },
        {
            "recipient": "Madame, Monsieur,",
            "subject": "Objet : Candidature",
            "intro": "Intro",
            "body_1": "Body 1",
            "body_2": "Body 2",
            "closing": "Closing",
            "signature": "Lucas Pertusa",
        },
        session_root,
    )

    updated = template_sessions.update_template_session(
        session["session_id"],
        session_root,
        {
            "cv": {
                "profile": "Profil mis à jour",
                "skills_text": "Sourcing\nSRM",
                "experience_0_bullets": "Nouveau bullet\nLead time",
            },
            "lm": {"body_1": "Argument modifié"},
        },
    )
    zip_path = template_sessions.export_template_session_zip(session["session_id"], session_root, tmp_path / "exports")

    root = session_root / session["session_id"]
    assert updated["cv_data"]["profile"] == "Profil mis à jour"
    assert updated["cv_data"]["skills"] == ["Sourcing", "SRM"]
    assert updated["cv_data"]["experiences"][0]["bullets"] == ["Nouveau bullet", "Lead time"]
    assert updated["lm_data"]["body_1"] == "Argument modifié"
    assert "Profil mis à jour" in "\n".join(paragraph.text for paragraph in Document(root / "CV_Lucas_Pertusa.docx").paragraphs)

    with zipfile.ZipFile(zip_path) as archive:
        assert sorted(archive.namelist()) == ["CV_Lucas_Pertusa.docx", "Lettre_Motivation_Lucas_Pertusa.docx"]
