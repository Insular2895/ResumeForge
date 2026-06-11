from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from src.web.experience_intake import add_validated_experience, propose_experience


def _master_profile(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "experiences"
    sheet.append(
        [
            "experience_id",
            "company",
            "location",
            "job_title",
            "date_start",
            "date_end",
            "industry_tags",
            "job_family_tags",
            "context",
            *[f"truth_bullet_{index}" for index in range(1, 11)],
            "tools_verified",
            "skills_verified",
            "skills_transferable",
            "skills_exposed",
            "kpis_verified",
            "evidence_strength",
            "allowed_rewrite_blocks",
            "cv_priority",
        ]
    )
    workbook.save(path)
    return path


def test_propose_experience_structures_five_factual_bullets():
    captured = {}

    def fake_generator(prompt):
        captured["prompt"] = prompt
        return """{
          "company": "Passy Primeur",
          "job_title": "Primeur sur les marchés",
          "location": "",
          "date_start": "",
          "date_end": "",
          "context": "Expérience terrain",
          "bullets": ["A", "B", "C", "D", "E"],
          "tools_verified": ["Bloomberg"],
          "skills_verified": ["supplier_negotiation"],
          "skills_transferable": ["pricing_analysis"],
          "industry_tags": ["fruits_legumes"],
          "job_family_tags": ["achats"]
        }"""

    proposal = propose_experience(
        "J'ai travaillé chez Passy Primeur et négocié à Rungis.",
        "Acheteur fruits et légumes",
        generator=fake_generator,
    )

    assert proposal["company"] == "Passy Primeur"
    assert proposal["bullets"] == ["A", "B", "C", "D", "E"]
    assert "n'invente aucun fait" in captured["prompt"]
    assert "Acheteur fruits et légumes" in captured["prompt"]


def test_proposal_keeps_unknown_identity_fields_empty_for_user_completion():
    proposal = propose_experience(
        "J'ai négocié des achats, organisé le transport et ajusté les prix en boutique.",
        generator=lambda prompt: """{
          "company": "",
          "job_title": "",
          "bullets": ["A", "B", "C", "D"],
          "tools_verified": [],
          "skills_verified": [],
          "skills_transferable": [],
          "industry_tags": [],
          "job_family_tags": []
        }""",
    )

    assert proposal["company"] == ""
    assert proposal["job_title"] == ""


def test_add_validated_experience_writes_locked_master_row_and_backup(tmp_path):
    master = _master_profile(tmp_path / "master_profile.xlsx")
    backup_dir = tmp_path / "backups"
    proposal = {
        "company": "Passy Primeur",
        "job_title": "Primeur sur les marchés",
        "location": "",
        "date_start": "",
        "date_end": "",
        "context": "Validé par l'utilisateur.",
        "bullets": ["Achat à Rungis.", "Négociation fournisseurs.", "Transport.", "Veille.", "Pricing."],
        "tools_verified": ["Bloomberg", "World Monitor"],
        "skills_verified": ["supplier_negotiation"],
        "skills_transferable": ["pricing_analysis"],
        "industry_tags": ["fruits_legumes"],
        "job_family_tags": ["achats"],
    }

    add_validated_experience(proposal, master_path=master, backup_dir=backup_dir)

    workbook = load_workbook(master, read_only=True, data_only=True)
    sheet = workbook["experiences"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    row = 2
    assert sheet.cell(row, headers["company"]).value == "Passy Primeur"
    assert sheet.cell(row, headers["truth_bullet_5"]).value == "Pricing."
    assert sheet.cell(row, headers["evidence_strength"]).value == "user_validated"
    assert sheet.cell(row, headers["tools_verified"]).value == "Bloomberg | World Monitor"
    assert list(backup_dir.glob("master_profile_before_experience_*.xlsx"))


def test_add_validated_experience_rejects_duplicate_company_and_role(tmp_path):
    master = _master_profile(tmp_path / "master_profile.xlsx")
    proposal = {
        "company": "Passy Primeur",
        "job_title": "Primeur sur les marchés",
        "bullets": ["A", "B", "C", "D"],
    }
    add_validated_experience(proposal, master_path=master, backup_dir=tmp_path / "backups")

    with pytest.raises(ValueError, match="existe déjà"):
        add_validated_experience(proposal, master_path=master, backup_dir=tmp_path / "backups")
