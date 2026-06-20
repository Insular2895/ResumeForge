from pathlib import Path
import json

from openpyxl import Workbook, load_workbook
import pytest

from src.application.experience_memory import (
    add_validated_memory,
    archive_memory,
    attach_memory_to_experiences,
    experience_options,
    profile_evidence_text,
)


def _master(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "experiences"
    sheet.append(["experience_id", "company", "job_title", "truth_bullet_1", "skills_exposed", "skills_verified", "context"])
    sheet.append(["exp_1", "Blurry", "Analyste ADV", "Gestion de partenaires et stocks.", "cybersécurité non vérifiée", "Excel", "oncologie non vérifiée"])
    workbook.save(path)
    return path


def test_validated_memory_is_appended_without_changing_truth_bullets(tmp_path):
    master = _master(tmp_path / "master.xlsx")

    add_validated_memory(
        {
            "experience_id": "exp_1",
            "target_domain": "data_analytics",
            "free_text": "Je construisais des tableaux Excel pour suivre les prix.",
            "qa_pairs": [{"question": "Créiez-vous des reportings ?", "answer": "Je produisais un reporting hebdomadaire."}],
        },
        master_path=master,
        backup_dir=tmp_path / "backups",
    )

    workbook = load_workbook(master, read_only=True, data_only=True)
    assert workbook["experiences"]["D2"].value == "Gestion de partenaires et stocks."
    memory = workbook["experience_memory"]
    assert memory["D2"].value == "Je construisais des tableaux Excel pour suivre les prix."
    assert memory["H2"].value == "user_validated"
    assert json.loads(memory["E2"].value)[0]["question"] == "Créiez-vous des reportings ?"
    assert list((tmp_path / "backups").glob("master_profile_before_memory_*.xlsx"))


def test_profile_evidence_includes_validated_memory(tmp_path):
    master = _master(tmp_path / "master.xlsx")
    add_validated_memory(
        {
            "experience_id": "exp_1",
            "target_domain": "data_analytics",
            "free_text": "Création de tableaux de bord Excel.",
            "qa_pairs": [],
        },
        master_path=master,
        backup_dir=tmp_path / "backups",
    )

    evidence = profile_evidence_text(master)

    assert "Gestion de partenaires et stocks" in evidence
    assert "Création de tableaux de bord Excel" in evidence
    assert "cybersécurité non vérifiée" not in evidence
    assert "oncologie non vérifiée" not in evidence


def test_memory_requires_a_real_experience_id_when_one_is_supplied(tmp_path):
    master = _master(tmp_path / "master.xlsx")

    with pytest.raises(ValueError, match="Expérience inconnue"):
        add_validated_memory(
            {"experience_id": "Blurry", "free_text": "Suivi Excel.", "qa_pairs": []},
            master_path=master,
            backup_dir=tmp_path / "backups",
        )


def test_memory_must_be_attached_to_a_real_experience(tmp_path):
    master = _master(tmp_path / "master.xlsx")

    with pytest.raises(ValueError, match="Rattache"):
        add_validated_memory(
            {"experience_id": "", "free_text": "Suivi Excel.", "qa_pairs": []},
            master_path=master,
            backup_dir=tmp_path / "backups",
        )


def test_identical_memory_is_deduplicated_and_backups_are_unique(tmp_path):
    master = _master(tmp_path / "master.xlsx")
    payload = {"experience_id": "exp_1", "free_text": "Suivi Excel.", "qa_pairs": []}

    add_validated_memory(payload, master_path=master, backup_dir=tmp_path / "backups")
    add_validated_memory(payload, master_path=master, backup_dir=tmp_path / "backups")

    workbook = load_workbook(master, read_only=True, data_only=True)
    assert workbook["experience_memory"].max_row == 2
    assert len(list((tmp_path / "backups").glob("master_profile_before_memory_*.xlsx"))) == 2


def test_linked_active_memory_is_attached_to_its_real_experience_only(tmp_path):
    master = _master(tmp_path / "master.xlsx")
    add_validated_memory(
        {"experience_id": "exp_1", "free_text": "Reporting hebdomadaire Excel.", "qa_pairs": []},
        master_path=master,
        backup_dir=tmp_path / "backups",
    )

    rows = attach_memory_to_experiences(master)

    assert rows[0]["experience_id"] == "exp_1"
    assert "Reporting hebdomadaire Excel" in rows[0]["validated_memory"]


def test_archived_memory_is_removed_from_evidence_and_experience_attachment(tmp_path):
    master = _master(tmp_path / "master.xlsx")
    add_validated_memory(
        {"experience_id": "exp_1", "free_text": "Ancienne information.", "qa_pairs": []},
        master_path=master,
        backup_dir=tmp_path / "backups",
    )
    memory_id = load_workbook(master, read_only=True, data_only=True)["experience_memory"]["A2"].value

    archive_memory(memory_id, master_path=master, backup_dir=tmp_path / "backups")

    assert "Ancienne information" not in profile_evidence_text(master)
    assert attach_memory_to_experiences(master)[0]["validated_memory"] == ""


def test_experience_options_expose_ids_with_readable_labels(tmp_path):
    master = _master(tmp_path / "master.xlsx")

    assert experience_options(master) == [{"id": "exp_1", "label": "Blurry · Analyste ADV"}]


def test_legacy_answers_json_remains_readable_after_schema_upgrade(tmp_path):
    master = _master(tmp_path / "master.xlsx")
    workbook = load_workbook(master)
    memory = workbook.create_sheet("experience_memory")
    memory.append(["memory_id", "experience_id", "target_domain", "free_text", "answers_json", "created_at", "source", "validation_status"])
    memory.append(["old_1", "exp_1", "data", "", '["Reporting Excel hebdomadaire."]', "", "web", "user_validated"])
    workbook.save(master)

    add_validated_memory(
        {"experience_id": "exp_1", "free_text": "Nouvelle preuve.", "qa_pairs": []},
        master_path=master,
        backup_dir=tmp_path / "backups",
    )

    assert "Reporting Excel hebdomadaire" in profile_evidence_text(master)
