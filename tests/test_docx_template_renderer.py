from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from src.render.docx_template import DocxTemplateRenderer


def test_renderer_removes_dangling_separators_for_missing_optional_values(tmp_path):
    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("[[POSITION]] - ")
    document.add_paragraph("[[LOCATION]] - [[DATES]]")
    document.save(template)

    output = tmp_path / "output.docx"
    DocxTemplateRenderer(template).render(
        {
            "[[POSITION]]": "Primeur sur les marchés",
            "[[LOCATION]]": "",
            "[[DATES]]": "",
        },
        output,
    )

    rendered = Document(output)
    lines = [paragraph.text.strip() for paragraph in rendered.paragraphs if paragraph.text.strip()]
    assert lines == ["Primeur sur les marchés"]


def test_renderer_removes_hidden_column_layout_that_creates_large_blanks(tmp_path):
    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("Expérience")
    break_paragraph = document.add_paragraph()
    break_paragraph.add_run().add_break(WD_BREAK.COLUMN)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "3")
    document.sections[0]._sectPr.append(cols)
    document.add_paragraph("[[BULLET]]")
    document.save(template)

    output = tmp_path / "output.docx"
    DocxTemplateRenderer(template).render({"[[BULLET]]": "Preuve métier."}, output)

    rendered = Document(output)
    assert all(
        br.get(qn("w:type")) != "column"
        for paragraph in rendered.paragraphs
        for br in paragraph._p.findall(".//" + qn("w:br"))
    )
    assert all(
        cols.get(qn("w:num")) in {None, "1"}
        for sect_pr in rendered._element.findall(".//" + qn("w:sectPr"))
        for cols in sect_pr.findall(qn("w:cols"))
    )


def test_renderer_removes_blank_paragraphs_before_experience_bullets(tmp_path):
    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("[[EXP_1_LOCATION]] - [[EXP_1_DATES]]")
    for _ in range(3):
        blank = document.add_paragraph("")
        blank.paragraph_format.space_before = Pt(12)
    document.add_paragraph("[[EXP_1_BULLETS]]")
    document.save(template)

    output = tmp_path / "output.docx"
    DocxTemplateRenderer(template).render(
        {
            "[[EXP_1_LOCATION]]": "Paris",
            "[[EXP_1_DATES]]": "2024 - 2025",
            "[[EXP_1_BULLETS]]": ["Coordination des flux."],
        },
        output,
    )

    rendered = Document(output)
    lines = [paragraph.text.strip() for paragraph in rendered.paragraphs if paragraph.text.strip()]
    assert lines == ["Paris - 2024 - 2025", "Coordination des flux."]


def test_renderer_adds_breathing_room_after_last_experience_bullet(tmp_path):
    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("[[EXP_1_BULLETS]]")
    document.add_paragraph("[[EXP_2_COMPAGNY]]")
    document.save(template)

    output = tmp_path / "output.docx"
    DocxTemplateRenderer(template).render(
        {
            "[[EXP_1_BULLETS]]": ["Premier bullet.", "Dernier bullet."],
            "[[EXP_2_COMPAGNY]]": "Weplugworld",
        },
        output,
    )

    rendered = Document(output)
    paragraphs = [paragraph for paragraph in rendered.paragraphs if paragraph.text.strip()]
    assert [paragraph.text for paragraph in paragraphs] == ["Premier bullet.", "Dernier bullet.", "Weplugworld"]
    assert paragraphs[0].paragraph_format.space_after.pt <= 2
    assert paragraphs[1].paragraph_format.space_after.pt >= 5
