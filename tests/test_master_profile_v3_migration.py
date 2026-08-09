from openpyxl import Workbook, load_workbook

from scripts.migrate_master_profile_v3 import SAP_EASY_ACCESS, SAP_EWM, migrate_workbook


def _legacy_workbook(path):
    workbook = Workbook()
    experiences = workbook.active
    experiences.title = "experiences"
    experiences.append(["experience_id", "company", "job_title", "context", "truth_bullet_1"])
    experiences.append(["normal_1", "Blurry", "Gestionnaire ADV", "alternance", "Suivi des commandes."])
    leadership = workbook.create_sheet("leadership")
    leadership.append(
        ["Role", "Organisation", "City", "Year", "truth_bullet_1", "validated_memory", "tools_verified"]
    )
    leadership.append(
        ["Gestionnaire flux export", "Minero", "Paris", "2020-2023", "Pilotage des expéditions.", "Fait validé", "Excel"]
    )
    certifications = workbook.create_sheet("certifications")
    certifications.append(["cert_name", "issuer", "status"])
    certifications.append(["SAP EASY ACCESS", "SAP", "verified"])
    certifications.append(["SAP Supply chain management EWM - 4HANA", "SAP", "verified"])
    workbook.save(path)


def test_migration_moves_leadership_to_freelance_without_duplication(tmp_path):
    path = tmp_path / "master.xlsx"
    _legacy_workbook(path)
    first = migrate_workbook(path, backup_dir=tmp_path / "backups")
    second = migrate_workbook(path, backup_dir=tmp_path / "backups")
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["experiences"]
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]

    assert first.migrated_experiences == 1
    assert second.changed is False
    assert len(rows) == 2
    migrated = next(row for row in rows if row["company"] == "Minero")
    assert migrated["job_title"] == "Gestionnaire flux export"
    assert migrated["dates"] == "2020-2023"
    assert migrated["truth_bullet_1"] == "Pilotage des expéditions."
    assert migrated["validated_memory"] == "Fait validé"
    assert migrated["is_freelance"] is True
    assert workbook["leadership_legacy"] is not None

    normal = next(row for row in rows if row["company"] == "Blurry")
    assert normal["is_freelance"] in (None, False)
    assert "Freelance" not in str(normal["job_title"])
    assert "Alternance" not in str(normal["job_title"])

    certification_names = [row[0] for row in workbook["certifications"].iter_rows(min_row=2, values_only=True)]
    assert SAP_EASY_ACCESS in certification_names
    assert SAP_EWM in certification_names
