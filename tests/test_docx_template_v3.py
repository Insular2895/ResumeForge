from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from src.generate_cv import build_replacements
from src.render.docx_template import DEFAULT_FONT_NAME, DocxTemplateRenderer


def _text(document):
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_v3_template_is_canonical_and_single_section():
    document = Document(Path("templates/base_cv_example.docx"))
    text = _text(document)
    assert "[[CV_HEADLINE]]" in text
    assert "[[EXP_3_COMPANY]]" in text
    assert "COMPAGNY" not in text
    assert "LEAD_1" not in text
    assert "Intérêts" not in text
    assert len(document.sections) == 1
    assert not document._element.findall(".//" + qn("w:pPr") + "/" + qn("w:sectPr"))
    assert len(document._element.findall(".//" + qn("w:drawing"))) == 1
    anchors = document._element.findall(".//" + qn("wp:anchor"))
    assert len(anchors) == 1
    assert anchors[0].get("behindDoc") == "0"
    assert anchors[0].find(qn("wp:wrapNone")) is not None
    assert anchors[0].find(qn("wp:wrapSquare")) is None
    education = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Éducation")
    assert education.paragraph_format.left_indent.pt == 0
    assert education.paragraph_format.right_indent.pt == 0
    assert education.paragraph_format.first_line_indent.pt == 0


def test_renderer_outputs_calibri_headline_third_experience_and_targeted_freelance(tmp_path):
    experiences = [
        {"company": "Blurry", "position": "Gestionnaire ADV", "display_position": "Gestionnaire ADV", "location": "Paris", "dates": "2024-2025", "bullets": ["Gestion des commandes."]},
        {"company": "Orion", "position": "Assistant commercial", "display_position": "Assistant commercial", "location": "Paris", "dates": "2024", "bullets": ["Suivi budgétaire."]},
        {"company": "Minero", "position": "Gestionnaire export", "display_position": "Gestionnaire export | Freelance", "location": "Paris", "dates": "2020-2023", "bullets": ["Pilotage des expéditions."]},
    ]
    replacements = build_replacements(experiences, ["SAP Easy Access"], ["SAP", "Excel"], "GESTIONNAIRE ADV EXPORT")
    output = tmp_path / "cv.docx"
    DocxTemplateRenderer("templates/base_cv_example.docx").render(replacements, output)
    document = Document(output)
    text = _text(document)

    assert DEFAULT_FONT_NAME == "Calibri"
    assert "GESTIONNAIRE ADV EXPORT" in text
    assert "Minero" in text
    assert "Gestionnaire export | Freelance" in text
    assert "Gestionnaire ADV |" not in text
    assert all(contract not in text for contract in ("Stage", "Alternance", "CDI", "CDD"))
    assert "[[" not in text
    assert "Intérêts" not in text
    assert "Leadership" not in text
    assert len(document._element.findall(".//" + qn("w:drawing"))) == 1
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                assert run.font.name == "Calibri"
