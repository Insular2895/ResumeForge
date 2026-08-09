"""Garde-fou de pagination du CV final.

LibreOffice fournit la mesure forte lorsqu'il est disponible. Le fallback
interne estime les lignes réellement rendues à partir des dimensions A4, des
tailles de police et du contenu du DOCX. Un dépassement non résolu bloque
l'export au lieu de laisser silencieusement un CV de deux pages.
"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import asdict, dataclass
import math
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION_START


class CvLayoutError(RuntimeError):
    pass


@dataclass(frozen=True)
class PageMeasurement:
    page_count: int
    method: str
    estimated_lines: float | None = None


@dataclass(frozen=True)
class LayoutGuardResult:
    output_path: str
    page_count: int
    measurement_method: str
    compaction_passes: int
    experiences: tuple[dict, ...]
    technical_skills: tuple[str, ...]

    def to_dict(self) -> dict:
        return asdict(self)


def _pdf_page_count(pdf_path: Path) -> int:
    pdfinfo = shutil.which("pdfinfo")
    if pdfinfo:
        completed = subprocess.run(
            [pdfinfo, str(pdf_path)], capture_output=True, text=True, timeout=30, check=False
        )
        match = re.search(r"^Pages:\s*(\d+)", completed.stdout, flags=re.MULTILINE)
        if match:
            return int(match.group(1))
    data = pdf_path.read_bytes()
    return len(re.findall(rb"/Type\s*/Page(?!s)\b", data))


def _measure_with_libreoffice(docx_path: Path) -> PageMeasurement | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory(prefix="resumeforge-pages-") as directory:
        completed = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", directory, str(docx_path)],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )
        pdf_path = Path(directory) / f"{docx_path.stem}.pdf"
        if completed.returncode != 0 or not pdf_path.exists():
            return None
        count = _pdf_page_count(pdf_path)
        return PageMeasurement(page_count=max(1, count), method="libreoffice_pdf")


def _paragraph_line_estimate(paragraph, usable_width_points: float) -> float:
    text = paragraph.text.strip()
    if not text:
        return 0.0
    sizes = [run.font.size.pt for run in paragraph.runs if run.text.strip() and run.font.size]
    size = max(sizes) if sizes else 10.0
    bold_factor = 1.05 if any(run.bold for run in paragraph.runs if run.text.strip()) else 1.0
    # Largeur moyenne Calibri ≈ 0,49 em. Le facteur protège contre une
    # estimation trop optimiste et les mots longs non sécables.
    chars_per_line = max(28.0, usable_width_points / (size * 0.54 * bold_factor))
    lines = max(1.0, math.ceil(len(text) / chars_per_line))
    before = paragraph.paragraph_format.space_before.pt if paragraph.paragraph_format.space_before else 0
    after = paragraph.paragraph_format.space_after.pt if paragraph.paragraph_format.space_after else 0
    return lines + (before + after) / max(9.0, size * 1.15)


def _measure_internally(docx_path: Path) -> PageMeasurement:
    document = Document(docx_path)
    section = document.sections[0]
    # Les soustractions python-docx retournent des EMU bruts (12 700 EMU/pt).
    usable_width = (section.page_width - section.left_margin - section.right_margin) / 12700
    usable_height = (section.page_height - section.top_margin - section.bottom_margin) / 12700
    lines_per_page = max(34.0, usable_height / 12.2)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    estimated_lines = sum(_paragraph_line_estimate(paragraph, usable_width) for paragraph in paragraphs)
    explicit_breaks = len(document._element.xpath(".//w:br[@w:type='page']"))
    section_breaks = sum(
        1
        for current in document.sections[1:]
        if current.start_type in {WD_SECTION_START.NEW_PAGE, WD_SECTION_START.ODD_PAGE, WD_SECTION_START.EVEN_PAGE}
    )
    page_count = max(1 + explicit_breaks + section_breaks, math.ceil(estimated_lines / lines_per_page))
    return PageMeasurement(page_count=page_count, method="internal_a4_estimate", estimated_lines=estimated_lines)


def measure_docx_pages(docx_path: str | Path) -> PageMeasurement:
    path = Path(docx_path)
    if not path.exists():
        raise FileNotFoundError(path)
    return _measure_with_libreoffice(path) or _measure_internally(path)


def _shorten_bullet(value: str, target: int) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    text = text.replace("dans le cadre de", "pour").replace("afin de", "pour")
    if len(text) <= target:
        return text
    # Suppression d'une proposition secondaire entière avant toute coupe.
    for separator in (", permettant ", ", contribuant ", ", incluant ", ", avec "):
        head = text.split(separator, 1)[0].rstrip(" ,;:")
        if 70 <= len(head) <= target:
            return head + ("." if not head.endswith(".") else "")
    boundary = max(text.rfind(";", 0, target), text.rfind(",", 0, target), text.rfind(" ", 0, target))
    if boundary < 70:
        return text
    shortened = text[:boundary].rstrip(" ,;:")
    return shortened + ("." if not shortened.endswith(".") else "")


def compact_cv_content(
    experiences: list[dict],
    technical_skills: list[str],
    pass_number: int,
) -> tuple[list[dict], list[str]]:
    compacted = deepcopy(experiences)
    limits = (4, 4, 2)
    target_length = 190 if pass_number == 1 else 155
    for index, experience in enumerate(compacted):
        bullet_limit = limits[index] if index < len(limits) else 0
        if pass_number >= 2 and index < 2:
            bullet_limit = 3
        experience["bullets"] = [
            _shorten_bullet(bullet, target_length)
            for bullet in experience.get("bullets", [])[:bullet_limit]
        ]

    skills = list(technical_skills)
    if skills:
        keep_ratio = 0.82 if pass_number == 1 else 0.62
        target_chars = max(70, int(sum(len(skill) + 3 for skill in skills) * keep_ratio))
        kept = []
        used = 0
        for skill in skills:
            cost = len(skill) + (3 if kept else 0)
            if kept and used + cost > target_chars:
                continue
            kept.append(skill)
            used += cost
        skills = kept
    return compacted, skills


def render_cv_one_page(
    *,
    renderer,
    output_path: str | Path,
    replacement_factory,
    experiences: list[dict],
    certifications: list[str],
    technical_skills: list[str],
    max_compaction_passes: int = 2,
) -> LayoutGuardResult:
    path = Path(output_path)
    current_experiences = deepcopy(experiences)
    current_skills = list(technical_skills)
    last_measurement = None
    for attempt in range(max_compaction_passes + 1):
        replacements = replacement_factory(current_experiences, certifications, current_skills)
        renderer.render(replacements, path)
        last_measurement = measure_docx_pages(path)
        if last_measurement.page_count == 1:
            return LayoutGuardResult(
                output_path=str(path),
                page_count=1,
                measurement_method=last_measurement.method,
                compaction_passes=attempt,
                experiences=tuple(current_experiences),
                technical_skills=tuple(current_skills),
            )
        if attempt < max_compaction_passes:
            current_experiences, current_skills = compact_cv_content(
                current_experiences, current_skills, attempt + 1
            )

    path.unlink(missing_ok=True)
    raise CvLayoutError(
        "Le CV dépasse encore une page après compaction prudente "
        f"({last_measurement.page_count if last_measurement else '?'} pages estimées)."
    )


__all__ = [
    "CvLayoutError",
    "LayoutGuardResult",
    "PageMeasurement",
    "compact_cv_content",
    "measure_docx_pages",
    "render_cv_one_page",
]
