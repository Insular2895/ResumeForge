from __future__ import annotations

from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook

from src.config import MASTER_PROFILE_PATH, OUTPUT_DIR


BACKUP_DIR = OUTPUT_DIR / "backups"
MEMORY_SHEET = "experience_memory"
MEMORY_HEADERS = [
    "memory_id",
    "experience_id",
    "target_domain",
    "free_text",
    "qa_pairs_json",
    "created_at",
    "source",
    "validation_status",
    "memory_status",
    "content_hash",
]
EVIDENCE_COLUMNS = {
    "experiences": {
        "company",
        "job_title",
        "position_title",
        "tools_verified",
        "skills_verified",
        "kpis_verified",
        *(f"truth_bullet_{index}" for index in range(1, 11)),
    },
    "leadership": {
        "organisation",
        "organization",
        "role",
        *(f"truth_bullet_{index}" for index in range(1, 11)),
    },
}


def _headers(sheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value not in (None, "")
    }


def _ensure_memory_sheet(workbook):
    if MEMORY_SHEET not in workbook.sheetnames:
        sheet = workbook.create_sheet(MEMORY_SHEET)
        sheet.append(MEMORY_HEADERS)
        return sheet
    sheet = workbook[MEMORY_SHEET]
    existing = _headers(sheet)
    for header in MEMORY_HEADERS:
        if header not in existing:
            sheet.cell(row=1, column=sheet.max_column + 1, value=header)
            existing[header] = sheet.max_column
    return sheet


def _backup(master: Path, backup_dir: Path) -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(master, backup_dir / f"master_profile_before_memory_{timestamp}.xlsx")


def _save_atomic(workbook, master: Path) -> None:
    with tempfile.NamedTemporaryFile(dir=master.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, master)
    finally:
        temporary_path.unlink(missing_ok=True)


def _qa_pairs(payload: dict) -> list[dict[str, str]]:
    pairs = []
    for item in payload.get("qa_pairs", []):
        if not isinstance(item, dict):
            continue
        question = str(item.get("question") or "").strip()
        answer = str(item.get("answer") or "").strip()
        if answer:
            pairs.append({"question": question, "answer": answer})
    if not pairs:
        pairs = [
            {"question": "", "answer": str(answer).strip()}
            for answer in payload.get("answers", [])
            if str(answer).strip()
        ]
    return pairs


def _experience_ids(workbook) -> set[str]:
    if "experiences" not in workbook.sheetnames:
        return set()
    sheet = workbook["experiences"]
    headers = _headers(sheet)
    column = headers.get("experience_id")
    if not column:
        return set()
    return {
        str(sheet.cell(row=row, column=column).value).strip()
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=column).value not in (None, "")
    }


def _content_hash(experience_id: str, target_domain: str, free_text: str, qa_pairs: list[dict]) -> str:
    content = json.dumps(
        {
            "experience_id": experience_id,
            "target_domain": target_domain,
            "free_text": free_text,
            "qa_pairs": qa_pairs,
        },
        ensure_ascii=False,
        sort_keys=True,
    )
    return hashlib.sha256(content.casefold().encode("utf-8")).hexdigest()


def add_validated_memory(payload: dict, *, master_path=MASTER_PROFILE_PATH, backup_dir=BACKUP_DIR) -> Path:
    free_text = str(payload.get("free_text") or "").strip()
    pairs = _qa_pairs(payload)
    if not free_text and not pairs:
        raise ValueError("Ajoute au moins une information avant validation.")

    master = Path(master_path)
    backups = Path(backup_dir)
    workbook = load_workbook(master)
    experience_id = str(payload.get("experience_id") or "").strip()
    if not experience_id:
        raise ValueError("Rattache cette information à une expérience réelle avant validation.")
    if experience_id not in _experience_ids(workbook):
        raise ValueError("Expérience inconnue : sélectionne une expérience réelle du profil.")

    sheet = _ensure_memory_sheet(workbook)
    headers = _headers(sheet)
    target_domain = str(payload.get("target_domain") or "").strip()
    content_hash = _content_hash(experience_id, target_domain, free_text, pairs)
    _backup(master, backups)

    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row=row, column=headers["content_hash"]).value == content_hash
            and str(sheet.cell(row=row, column=headers["memory_status"]).value or "active") == "active"
        ):
            return master

    now = datetime.now()
    values = {
        "memory_id": f"memory_{now.strftime('%Y%m%d_%H%M%S_%f')}",
        "experience_id": experience_id,
        "target_domain": target_domain,
        "free_text": free_text,
        "qa_pairs_json": json.dumps(pairs, ensure_ascii=False),
        "created_at": now.isoformat(timespec="seconds"),
        "source": "web_enrichment",
        "validation_status": "user_validated",
        "memory_status": "active",
        "content_hash": content_hash,
    }
    sheet.append([values.get(header, "") for header in _headers(sheet)])
    _save_atomic(workbook, master)
    return master


def archive_memory(memory_id: str, *, master_path=MASTER_PROFILE_PATH, backup_dir=BACKUP_DIR) -> Path:
    master = Path(master_path)
    workbook = load_workbook(master)
    if MEMORY_SHEET not in workbook.sheetnames:
        raise ValueError("Mémoire introuvable.")
    sheet = _ensure_memory_sheet(workbook)
    headers = _headers(sheet)
    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row=row, column=headers["memory_id"]).value or "") == str(memory_id):
            _backup(master, Path(backup_dir))
            sheet.cell(row=row, column=headers["memory_status"], value="archived")
            _save_atomic(workbook, master)
            return master
    raise ValueError("Mémoire introuvable.")


def _memory_texts(sheet) -> list[tuple[str, str]]:
    headers = _headers(sheet)
    records = []
    for row in range(2, sheet.max_row + 1):
        status = str(sheet.cell(row=row, column=headers.get("validation_status", 1)).value or "")
        memory_status = str(sheet.cell(row=row, column=headers.get("memory_status", 1)).value or "active")
        if status != "user_validated" or memory_status == "archived":
            continue
        experience_id = str(sheet.cell(row=row, column=headers.get("experience_id", 1)).value or "").strip()
        if not experience_id:
            continue
        texts = [str(sheet.cell(row=row, column=headers.get("free_text", 1)).value or "").strip()]
        qa_header = headers.get("qa_pairs_json")
        if not qa_header or not sheet.cell(row=row, column=qa_header).value:
            qa_header = headers.get("answers_json")
        if qa_header:
            try:
                items = json.loads(str(sheet.cell(row=row, column=qa_header).value or "[]"))
                texts.extend(
                    str(item.get("answer") if isinstance(item, dict) else item).strip()
                    for item in items
                    if str(item.get("answer") if isinstance(item, dict) else item).strip()
                )
            except json.JSONDecodeError:
                pass
        records.append(
            (
                experience_id,
                "\n".join(text for text in texts if text),
            )
        )
    return records


def attach_memory_to_experiences(master_path=MASTER_PROFILE_PATH) -> list[dict]:
    workbook = load_workbook(Path(master_path), read_only=True, data_only=True)
    if "experiences" not in workbook.sheetnames:
        return []
    sheet = workbook["experiences"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    linked: dict[str, list[str]] = {}
    if MEMORY_SHEET in workbook.sheetnames:
        for experience_id, text in _memory_texts(workbook[MEMORY_SHEET]):
            if experience_id and text:
                linked.setdefault(experience_id, []).append(text)
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        row["validated_memory"] = "\n".join(linked.get(str(row.get("experience_id") or "").strip(), []))
        rows.append(row)
    return rows


def experience_options(master_path=MASTER_PROFILE_PATH) -> list[dict[str, str]]:
    master = Path(master_path)
    if not master.exists():
        return []
    workbook = load_workbook(master, read_only=True, data_only=True)
    if "experiences" not in workbook.sheetnames:
        return []
    sheet = workbook["experiences"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    options = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        experience_id = str(row.get("experience_id") or "").strip()
        if not experience_id:
            continue
        company = str(row.get("company") or row.get("organisation") or "").strip()
        title = str(row.get("job_title") or row.get("position_title") or row.get("role") or "").strip()
        label = " · ".join(value for value in [company, title] if value) or experience_id
        options.append({"id": experience_id, "label": label})
    return options


def profile_evidence_text(master_path=MASTER_PROFILE_PATH) -> str:
    master = Path(master_path)
    if not master.exists():
        return ""
    workbook = load_workbook(master, read_only=True, data_only=True)
    values = []
    for sheet_name, allowed in EVIDENCE_COLUMNS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        indexes = [index for index, header in enumerate(headers) if header in allowed]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values.extend(str(row[index]).strip() for index in indexes if row[index] not in (None, ""))
    if MEMORY_SHEET in workbook.sheetnames:
        values.extend(text for _, text in _memory_texts(workbook[MEMORY_SHEET]) if text)
    return "\n".join(values)
