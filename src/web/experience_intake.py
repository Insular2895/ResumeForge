from __future__ import annotations

from datetime import datetime
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata

from openpyxl import load_workbook

from src.config import MASTER_PROFILE_PATH, OUTPUT_DIR
from src.llm.gemini_client import ask_gemini


BACKUP_DIR = OUTPUT_DIR / "backups"
LIST_FIELDS = {
    "tools_verified",
    "skills_verified",
    "skills_transferable",
    "industry_tags",
    "job_family_tags",
}


def propose_experience(experience_text: str, job_text: str = "", *, generator=ask_gemini) -> dict:
    source = str(experience_text or "").strip()
    if len(source) < 30:
        raise ValueError("Décris davantage l'expérience avant de lancer l'analyse.")

    prompt = f"""
Transforme l'ancienne expérience fournie par l'utilisateur en fiche CV professionnelle à valider.

Règles absolues :
- n'invente aucun fait, chiffre, date, outil, employeur ou résultat ;
- utilise uniquement les faits explicitement présents dans le récit utilisateur ;
- chaque élément descriptif, bénéfice et résultat doit être directement traçable au récit utilisateur ;
- n'ajoute pas de qualificatif comme "concurrentiel", "stratégique", "quotidien" ou de bénéfice commercial implicite ;
- l'offre cible sert seulement à choisir le vocabulaire et l'ordre des preuves ;
- produis 4 à 5 bullets professionnels selon une logique ABC ou XYZ ;
- chaque bullet relie autant que possible une action, un contexte et un bénéfice prouvé ;
- laisse les champs inconnus vides ;
- utilise des formulations prudentes comme "contribution à" ou "participation à" si le niveau de responsabilité est ambigu ;
- retourne uniquement un objet JSON valide.

Schéma JSON :
{{
  "company": "",
  "job_title": "",
  "location": "",
  "date_start": "",
  "date_end": "",
  "context": "",
  "bullets": ["", "", "", "", ""],
  "tools_verified": [],
  "skills_verified": [],
  "skills_transferable": [],
  "industry_tags": [],
  "job_family_tags": []
}}

Récit utilisateur :
{source}

Offre cible facultative :
{str(job_text or "").strip()}
"""
    try:
        payload = json.loads(_clean_json(generator(prompt)))
    except (json.JSONDecodeError, TypeError) as exc:
        raise ValueError("La proposition d'expérience produite est illisible. Relance l'analyse.") from exc
    return normalize_proposal(payload)


def normalize_proposal(payload: dict, *, require_identity: bool = False) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("La proposition d'expérience doit être un objet structuré.")

    proposal = {
        key: str(payload.get(key, "") or "").strip()
        for key in ["company", "job_title", "location", "date_start", "date_end", "context"]
    }
    bullets = [str(item or "").strip() for item in payload.get("bullets", []) if str(item or "").strip()]
    if require_identity and (not proposal["company"] or not proposal["job_title"]):
        raise ValueError("Entreprise et intitulé de poste requis avant validation.")
    if not 4 <= len(bullets) <= 5:
        raise ValueError("La proposition doit contenir 4 à 5 bullets factuels.")
    proposal["bullets"] = bullets
    for field in LIST_FIELDS:
        proposal[field] = _list_value(payload.get(field, []))
    return proposal


def add_validated_experience(
    proposal: dict,
    *,
    master_path: str | Path = MASTER_PROFILE_PATH,
    backup_dir: str | Path = BACKUP_DIR,
) -> Path:
    normalized = normalize_proposal(proposal, require_identity=True)
    master = Path(master_path)
    backups = Path(backup_dir)
    if not master.exists():
        raise ValueError("Le profil Excel maître est introuvable.")

    workbook = load_workbook(master)
    if "experiences" not in workbook.sheetnames:
        raise ValueError("La feuille experiences est absente du profil maître.")
    sheet = workbook["experiences"]
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}

    required = {
        "experience_id",
        "company",
        "job_title",
        "evidence_strength",
        *[f"truth_bullet_{index}" for index in range(1, 6)],
    }
    missing = sorted(required - headers.keys())
    if missing:
        raise ValueError(f"Colonne requise absente du profil maître : {missing[0]}")

    company_key = normalized["company"].casefold()
    role_key = normalized["job_title"].casefold()
    for row in range(2, sheet.max_row + 1):
        company = str(sheet.cell(row, headers["company"]).value or "").strip().casefold()
        role = str(sheet.cell(row, headers["job_title"]).value or "").strip().casefold()
        if company == company_key and role == role_key:
            raise ValueError("Cette expérience existe déjà dans le profil maître.")

    backups.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(master, backups / f"master_profile_before_experience_{timestamp}.xlsx")

    row = sheet.max_row + 1
    values = {
        "experience_id": f"{_slug(normalized['company'])}_{_slug(normalized['job_title'])}_{timestamp}",
        "company": normalized["company"],
        "location": normalized["location"],
        "job_title": normalized["job_title"],
        "date_start": normalized["date_start"],
        "date_end": normalized["date_end"],
        "industry_tags": _pipe(normalized["industry_tags"]),
        "job_family_tags": _pipe(normalized["job_family_tags"]),
        "context": normalized["context"] or "Expérience confirmée et validée par l'utilisateur.",
        "tools_verified": _pipe(normalized["tools_verified"]),
        "skills_verified": _pipe(normalized["skills_verified"]),
        "skills_transferable": _pipe(normalized["skills_transferable"]),
        "evidence_strength": "user_validated",
        "allowed_rewrite_blocks": "locked_user_validated_bullets",
        "cv_priority": 5,
    }
    for index, bullet in enumerate(normalized["bullets"], start=1):
        values[f"truth_bullet_{index}"] = bullet
    for key, value in values.items():
        if key in headers:
            sheet.cell(row, headers[key], value)

    with tempfile.NamedTemporaryFile(dir=master.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, master)
    finally:
        temporary_path.unlink(missing_ok=True)
    return master


def _clean_json(raw: str) -> str:
    text = str(raw or "").strip()
    if text.startswith("```json"):
        text = text[7:].strip()
    elif text.startswith("```"):
        text = text[3:].strip()
    if text.endswith("```"):
        text = text[:-3].strip()
    return text


def _list_value(value) -> list[str]:
    if isinstance(value, str):
        parts = re.split(r"[|;\n]+", value)
    elif isinstance(value, (list, tuple, set)):
        parts = value
    else:
        parts = []
    return [str(item).strip() for item in parts if str(item).strip()]


def _pipe(values) -> str:
    return " | ".join(_list_value(values))


def _slug(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char)).casefold()
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", text)).strip("_") or "experience"
