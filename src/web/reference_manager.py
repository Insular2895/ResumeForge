from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import os
import shutil
import tempfile

from docx import Document
from openpyxl import load_workbook


ROOT_DIR = Path(__file__).resolve().parents[2]


@dataclass(frozen=True)
class ReferenceSpec:
    key: str
    label: str
    destination: Path
    extensions: tuple[str, ...]
    required_placeholders: tuple[str, ...] = ()


REFERENCE_SPECS = {
    "master_profile": ReferenceSpec(
        key="master_profile",
        label="Profil Excel maître",
        destination=ROOT_DIR / "data" / "reference" / "master_profile.xlsx",
        extensions=(".xlsx",),
    ),
    "cv_template": ReferenceSpec(
        key="cv_template",
        label="Template CV Word",
        destination=ROOT_DIR / "templates" / "base_cv.docx",
        extensions=(".docx",),
        required_placeholders=("[[EXP_1_BULLETS]]",),
    ),
    "lm_template": ReferenceSpec(
        key="lm_template",
        label="Template LM Word",
        destination=ROOT_DIR / "templates" / "base_cover_letter.docx",
        extensions=(".docx",),
        required_placeholders=("[[LM_FINAL_LETTER]]",),
    ),
    "reference_cv": ReferenceSpec(
        key="reference_cv",
        label="CV de référence",
        destination=ROOT_DIR / "data" / "input" / "reference_cv.docx",
        extensions=(".docx", ".md", ".txt"),
    ),
}


def _spec(key: str) -> ReferenceSpec:
    try:
        return REFERENCE_SPECS[key]
    except KeyError as exc:
        raise ValueError(f"Référence inconnue : {key}") from exc


def _existing_reference_path(spec: ReferenceSpec) -> Path | None:
    if spec.key != "reference_cv":
        return spec.destination if spec.destination.exists() else None
    for extension in spec.extensions:
        candidate = spec.destination.with_suffix(extension)
        if candidate.exists():
            return candidate
    return None


def get_reference_statuses() -> dict[str, dict]:
    statuses: dict[str, dict] = {}
    for key, spec in REFERENCE_SPECS.items():
        path = _existing_reference_path(spec)
        error = ""
        ready = path is not None
        if path is not None:
            try:
                validate_reference(key, path)
            except ValueError as exc:
                ready = False
                error = str(exc)
        statuses[key] = {
            "key": key,
            "label": spec.label,
            "ready": ready,
            "filename": path.name if path else "",
            "error": error,
        }
    return statuses


def _document_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(chunks)


def validate_reference(key: str, source: str | Path) -> Path:
    spec = _spec(key)
    path = Path(source)
    if not path.exists() or not path.is_file():
        raise ValueError("Fichier introuvable.")
    if path.suffix.lower() not in spec.extensions:
        allowed = ", ".join(spec.extensions)
        raise ValueError(f"Extension invalide. Formats acceptés : {allowed}")
    if path.stat().st_size == 0:
        raise ValueError("Le fichier est vide.")

    if path.suffix.lower() == ".docx":
        try:
            text = _document_text(path)
        except Exception as exc:
            raise ValueError("Le document Word est invalide.") from exc
        missing = [placeholder for placeholder in spec.required_placeholders if placeholder not in text]
        if missing:
            raise ValueError(f"Template invalide : placeholder requis absent ({missing[0]}).")

    if path.suffix.lower() == ".xlsx":
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            workbook.close()
        except Exception as exc:
            raise ValueError("Le profil Excel est invalide.") from exc

    return path


def replace_reference(key: str, source: str | Path) -> Path:
    spec = _spec(key)
    source_path = validate_reference(key, source)
    destination = spec.destination
    if key == "reference_cv":
        destination = destination.with_suffix(source_path.suffix.lower())

    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        shutil.copy2(source_path, temporary_path)
        os.replace(temporary_path, destination)
    finally:
        temporary_path.unlink(missing_ok=True)

    if key == "reference_cv":
        for extension in spec.extensions:
            candidate = spec.destination.with_suffix(extension)
            if candidate != destination:
                candidate.unlink(missing_ok=True)
    return destination
