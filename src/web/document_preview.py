from __future__ import annotations

import base64
from datetime import datetime
from html.parser import HTMLParser
import html
import json
from pathlib import Path
import re
import uuid
import zipfile

from docx import Document

CV_SECTION_HEADINGS = {
    "éducation",
    "education",
    "expériences",
    "experiences",
    "leadership et activités",
    "leadership et activites",
    "compétences et intérêts",
    "competences et interets",
}

CV_INLINE_LABELS = {"certifications"}
CV_PREFIX_LABELS = ("Compétences techniques :", "Intérêts :", "Langues :")
CV_ENTRY_SECTIONS = {"expériences", "experiences", "leadership et activités", "leadership et activites"}


def _normalized_heading(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "").strip().lower())


def _is_cv_section_heading(text: str) -> bool:
    return _normalized_heading(text) in CV_SECTION_HEADINGS


def _first_image_data_uri(document: Document) -> str:
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        content_type = rel.target_part.content_type or "image/png"
        encoded = base64.b64encode(rel.target_part.blob).decode("ascii")
        return f"data:{content_type};base64,{encoded}"
    return ""


def _is_list_paragraph(paragraph) -> bool:
    text = paragraph.text.strip()
    if text.startswith(("•", "-", "*")):
        return True
    style = (paragraph.style.name or "").lower() if paragraph.style else ""
    if style.startswith("list") or "bullet" in style or "puce" in style:
        return True
    ppr = paragraph._p.pPr
    return bool(ppr is not None and ppr.numPr is not None)


def _paragraph_content_html(text: str) -> str:
    escaped = html.escape(text)
    if _normalized_heading(text) in CV_INLINE_LABELS:
        return f"<strong>{escaped}</strong>"
    for prefix in CV_PREFIX_LABELS:
        if text.startswith(prefix):
            suffix = html.escape(text[len(prefix) :].lstrip())
            return f"<strong>{html.escape(prefix)}</strong> {suffix}".rstrip()
    return escaped


class _PlainTextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._in_li = False

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"p", "h1", "h2", "h3", "h4", "div"}:
            self._newline(block=True)
        if tag == "li":
            self._newline(block=False)
            self.parts.append("- ")
            self._in_li = True
        if tag == "br":
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"p", "h1", "h2", "h3", "h4", "div", "ul", "ol"}:
            self._newline(block=True)
        if tag == "li":
            self._in_li = False
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        text = re.sub(r"\s+", " ", data).strip()
        if text:
            self.parts.append(text)

    def _newline(self, *, block: bool) -> None:
        current = "".join(self.parts)
        if not current:
            return
        suffix = "\n\n" if block else "\n"
        if not current.endswith(suffix):
            self.parts.append(suffix)

    def text(self) -> str:
        text = "".join(self.parts)
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()


def html_to_plain_text(html_content: str) -> str:
    parser = _PlainTextHTMLParser()
    parser.feed(str(html_content or ""))
    return parser.text()


class _DocumentHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.blocks: list[tuple[str, str]] = []
        self._stack: list[str] = []
        self._current_tag = ""
        self._current_text: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"h1", "h2", "h3", "p", "li"}:
            self._flush()
            self._current_tag = tag
            self._current_text = []
        self._stack.append(tag)

    def handle_endtag(self, tag: str) -> None:
        if tag == self._current_tag:
            self._flush()
        if self._stack:
            self._stack.pop()

    def handle_data(self, data: str) -> None:
        if not self._current_tag:
            return
        text = re.sub(r"\s+", " ", data).strip()
        if text:
            self._current_text.append(text)

    def _flush(self) -> None:
        if not self._current_tag:
            return
        text = " ".join(self._current_text).strip()
        if text:
            self.blocks.append((self._current_tag, html.unescape(text)))
        self._current_tag = ""
        self._current_text = []

    def close(self) -> None:
        self._flush()
        super().close()


def html_to_document_blocks(html_content: str) -> list[tuple[str, str]]:
    parser = _DocumentHTMLParser()
    parser.feed(str(html_content or ""))
    parser.close()
    return parser.blocks


class _EditableDocumentParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.image_html = ""
        self.blocks: list[dict] = []
        self._current_tag = ""
        self._current_text: list[str] = []
        self._strong_depth = 0
        self._saw_strong = False

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag == "img" and not self.image_html:
            attr_text = " ".join(
                f'{name}="{html.escape(str(value), quote=True)}"' for name, value in attrs if value is not None
            )
            self.image_html = f"<img {attr_text}>" if attr_text else "<img>"
            return
        if tag in {"h2", "p", "li"}:
            self._flush()
            self._current_tag = tag
            self._current_text = []
            self._saw_strong = False
        if tag in {"strong", "b"} and self._current_tag:
            self._strong_depth += 1
            self._saw_strong = True

    def handle_endtag(self, tag: str) -> None:
        if tag in {"strong", "b"} and self._strong_depth:
            self._strong_depth -= 1
        if tag == self._current_tag:
            self._flush()

    def handle_data(self, data: str) -> None:
        if not self._current_tag:
            return
        text = re.sub(r"\s+", " ", data).strip()
        if text:
            self._current_text.append(text)

    def _flush(self) -> None:
        if not self._current_tag:
            return
        text = " ".join(self._current_text).strip()
        if text:
            self.blocks.append(
                {
                    "id": f"block_{len(self.blocks)}",
                    "tag": self._current_tag,
                    "text": html.unescape(text),
                    "strong": self._saw_strong,
                }
            )
        self._current_tag = ""
        self._current_text = []
        self._saw_strong = False

    def close(self) -> None:
        self._flush()
        super().close()


def extract_editable_blocks(html_content: str) -> dict:
    parser = _EditableDocumentParser()
    parser.feed(str(html_content or ""))
    parser.close()
    return {"image_html": parser.image_html, "blocks": parser.blocks}


def render_editable_blocks(blocks: list[dict], *, image_html: str = "") -> str:
    chunks: list[str] = []
    if image_html:
        chunks.append(str(image_html))
    for block in blocks:
        tag = str(block.get("tag") or "p")
        text = str(block.get("text") or "")
        escaped = html.escape(text)
        if tag == "h2":
            chunks.append(f"<h2>{escaped}</h2>")
        elif tag == "li":
            chunks.append(f"<ul><li>{escaped}</li></ul>")
        elif block.get("strong"):
            chunks.append(f"<p>{_paragraph_content_html(text) if any(text.startswith(prefix) for prefix in CV_PREFIX_LABELS) else f'<strong>{escaped}</strong>'}</p>")
        else:
            chunks.append(f"<p>{escaped}</p>")
    return "\n".join(chunks).strip() or "<p>Document vide.</p>"


def _docx_to_html(path: Path) -> str:
    document = Document(path)
    chunks: list[str] = []
    image_data_uri = _first_image_data_uri(document)
    if image_data_uri:
        chunks.append(f'<img src="{image_data_uri}" alt="Photo de profil">')
    current_section = ""
    entry_line_index = 0
    just_finished_list = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if _is_cv_section_heading(text):
            current_section = _normalized_heading(text)
            entry_line_index = 0
            just_finished_list = False
            chunks.append(f"<h2>{html.escape(text)}</h2>")
        elif "heading" in style or "titre" in style:
            chunks.append(f"<p><strong>{html.escape(text)}</strong></p>")
        elif _is_list_paragraph(paragraph):
            chunks.append(f"<ul><li>{html.escape(text.lstrip('•-* ').strip())}</li></ul>")
            just_finished_list = True
        else:
            if current_section in CV_ENTRY_SECTIONS and just_finished_list:
                entry_line_index = 0
                just_finished_list = False
            content = _paragraph_content_html(text)
            if current_section in CV_ENTRY_SECTIONS and entry_line_index in {0, 1}:
                content = f"<strong>{html.escape(text)}</strong>"
            chunks.append(f"<p>{content}</p>")
            entry_line_index += 1
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(f"<p>{html.escape(' | '.join(cells))}</p>")
    return "\n".join(chunks).strip() or "<p>Document vide.</p>"


def _first_file(pack_dir: Path, pattern: str) -> Path | None:
    matches = sorted(path for path in pack_dir.glob(pattern) if path.is_file())
    return matches[0] if matches else None


def create_preview_session(pack_dir: str | Path, preview_root: str | Path, *, metadata: dict | None = None) -> dict:
    pack = Path(pack_dir)
    if not pack.is_dir():
        raise ValueError("Pack de candidature introuvable.")
    cv_path = _first_file(pack, "CV*.docx")
    lm_path = _first_file(pack, "LM*.docx")
    if not cv_path and not lm_path:
        raise ValueError("Aucun document éditable trouvé dans le pack.")

    preview_id = uuid.uuid4().hex
    root = Path(preview_root) / preview_id
    root.mkdir(parents=True, exist_ok=True)
    session = {
        "preview_id": preview_id,
        "pack_dir": str(pack),
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "cv_generated": _docx_to_html(cv_path) if cv_path else "<p>CV non généré.</p>",
        "cv_edited": "",
        "cv_is_dirty": False,
        "lm_generated": _docx_to_html(lm_path) if lm_path else "<p>Lettre de motivation non générée.</p>",
        "lm_edited": "",
        "lm_is_dirty": False,
        "metadata": metadata or {},
    }
    session["cv_document"] = extract_editable_blocks(session["cv_generated"])
    session["lm_document"] = extract_editable_blocks(session["lm_generated"])
    (root / "session.json").write_text(json.dumps(session, ensure_ascii=False, indent=2), encoding="utf-8")
    return session


def load_preview_session(preview_id: str, preview_root: str | Path) -> dict:
    path = Path(preview_root) / str(preview_id) / "session.json"
    if not path.exists():
        raise ValueError("Session de prévisualisation introuvable.")
    session = json.loads(path.read_text(encoding="utf-8"))
    changed = _backfill_missing_cv_photo(session)
    if _backfill_structured_documents(session):
        changed = True
    if changed:
        path.write_text(json.dumps(session, ensure_ascii=False, indent=2), encoding="utf-8")
    return session


def _backfill_missing_cv_photo(session: dict) -> bool:
    current_html = str(session.get("cv_generated") or "")
    if "<img" in current_html:
        return False
    pack_dir = session.get("pack_dir")
    if not pack_dir:
        return False
    cv_path = _first_file(Path(pack_dir), "CV*.docx")
    if not cv_path:
        return False
    refreshed_html = _docx_to_html(cv_path)
    if "<img" not in refreshed_html:
        return False
    session["cv_generated"] = refreshed_html
    session["cv_document"] = extract_editable_blocks(refreshed_html)
    if not session.get("cv_is_dirty"):
        session["cv_edited"] = ""
    return True


def _backfill_structured_documents(session: dict) -> bool:
    changed = False
    if not isinstance(session.get("cv_document"), dict):
        session["cv_document"] = extract_editable_blocks(str(session.get("cv_generated") or ""))
        changed = True
    if not isinstance(session.get("lm_document"), dict):
        session["lm_document"] = extract_editable_blocks(str(session.get("lm_generated") or ""))
        changed = True
    return changed


def _pdf_escape(text: str) -> str:
    escaped = text.encode("latin-1", errors="backslashreplace").decode("latin-1")
    return escaped.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


PDF_PAGE_WIDTH = 595
PDF_PAGE_HEIGHT = 842
PDF_MARGIN_X = 45
PDF_MARGIN_TOP = 51
PDF_MARGIN_BOTTOM = 45
PDF_BODY_SIZE = 10.6
PDF_BODY_LEADING = 13.6
PDF_HEADER_SIZE = 17.5
PDF_SECTION_SIZE = 11.5

DOCUMENT_EXPORT_CSS = """
* {
  box-sizing: border-box;
}

html,
body {
  margin: 0;
  background: #eef1f6;
  color: #111;
  font-family: Arial, Helvetica, sans-serif;
}

.document-workspace {
  min-height: 100vh;
  padding: 34px 24px 42px;
  background: #eef1f6;
}

.document-page {
  width: 210mm;
  min-height: 297mm;
  margin: 0 auto;
  padding: 17mm 12.5mm 11mm;
  border: 1px solid #d9dee8;
  background: #fff;
  box-shadow: 0 18px 46px rgb(23 32 51 / 14%);
}

.document-body {
  position: relative;
  min-height: calc(297mm - 28mm);
  color: #111;
  font-family: Arial, Helvetica, sans-serif;
  font-size: 11pt;
  line-height: 1.18;
}

.document-body h1,
.document-body h2,
.document-body h3 {
  color: #111;
  line-height: 1.12;
}

.document-body h2 {
  margin: 15px 0 12px;
  padding-bottom: 0;
  border-bottom: 0;
  font-size: 11.5pt;
  font-weight: 800;
  letter-spacing: 0;
  text-align: center;
  text-transform: none;
}

.document-body h3 {
  font-size: 10.8pt;
  font-weight: 800;
}

.document-body p {
  margin: 0 0 2px;
  color: #111;
}

.document-body ul {
  margin: 9px 0 14px;
  padding-left: 29px;
}

.document-body li {
  margin: 0 0 4px;
  padding-left: 7px;
}

.document-body img:first-child {
  position: absolute;
  top: 0;
  left: 0;
  width: 27mm;
  height: 27mm;
  border-radius: 999px;
  object-fit: cover;
}

.document-body p:first-of-type {
  text-align: center;
}

.document-body img:first-child + p {
  min-height: 27mm;
  margin: 0 0 4px;
  padding-left: 34mm;
  padding-top: 11mm;
  font-size: 10.7pt;
  line-height: 1.15;
}

.document-body p:first-of-type strong,
.document-body p:first-of-type b {
  font-size: 12pt;
}

.document-body h2 + p strong,
.document-body h2 + p b,
.document-body p:has(strong) {
  font-weight: 800;
}

@page {
  size: A4;
  margin: 0;
}

@media print {
  html,
  body,
  .document-workspace {
    background: #fff;
    padding: 0;
  }

  .document-page {
    width: 210mm;
    min-height: 297mm;
    margin: 0;
    border: 0;
    box-shadow: none;
  }
}
""".strip()


def _wrap_text(text: str, *, max_width: float, font_size: float) -> list[str]:
    words = str(text or "").split()
    if not words:
        return []
    average_char_width = font_size * 0.52
    max_chars = max(18, int(max_width / average_char_width))
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def _text_width_estimate(text: str, font_size: float) -> float:
    return len(text) * font_size * 0.52


def _render_pdf_commands(html_content: str) -> list[str]:
    blocks = html_to_document_blocks(html_content)
    if not blocks:
        blocks = [("p", "Document final.")]

    max_width = PDF_PAGE_WIDTH - (PDF_MARGIN_X * 2)
    y = PDF_PAGE_HEIGHT - PDF_MARGIN_TOP
    commands: list[str] = []
    body_index = 0

    def ensure_space(height: float) -> None:
        nonlocal y, commands
        if y - height < PDF_MARGIN_BOTTOM:
            commands.append("__PAGE_BREAK__")
            y = PDF_PAGE_HEIGHT - PDF_MARGIN_TOP

    def draw_text(line: str, *, x: float, size: float, font: str = "F1") -> None:
        commands.append("BT")
        commands.append(f"/{font} {size:.1f} Tf")
        commands.append(f"1 0 0 1 {x:.1f} {y:.1f} Tm")
        commands.append(f"({_pdf_escape(line)}) Tj")
        commands.append("ET")

    for tag, raw_text in blocks:
        text = raw_text.strip()
        if not text:
            continue

        if tag == "h1" or (tag == "p" and body_index == 0):
            lines = _wrap_text(text, max_width=max_width, font_size=PDF_HEADER_SIZE)
            ensure_space(len(lines) * 20 + 6)
            for line in lines:
                x = max(PDF_MARGIN_X, (PDF_PAGE_WIDTH - _text_width_estimate(line, PDF_HEADER_SIZE)) / 2)
                draw_text(line, x=x, size=PDF_HEADER_SIZE, font="F2")
                y -= 20
            y -= 1
            body_index += 1
            continue

        if tag == "p" and body_index == 1:
            lines = _wrap_text(text, max_width=max_width, font_size=9.5)
            ensure_space(len(lines) * 12 + 9)
            for line in lines:
                x = max(PDF_MARGIN_X, (PDF_PAGE_WIDTH - _text_width_estimate(line, 9.5)) / 2)
                draw_text(line, x=x, size=9.5, font="F1")
                y -= 12
            y -= 7
            body_index += 1
            continue

        if tag == "h2":
            section = text.upper()
            ensure_space(24)
            y -= 4
            draw_text(section, x=PDF_MARGIN_X, size=PDF_SECTION_SIZE, font="F2")
            y -= 4
            commands.append(f"{PDF_MARGIN_X:.1f} {y:.1f} m {PDF_PAGE_WIDTH - PDF_MARGIN_X:.1f} {y:.1f} l S")
            y -= 12
            body_index += 1
            continue

        if tag == "h3":
            lines = _wrap_text(text, max_width=max_width, font_size=PDF_BODY_SIZE)
            ensure_space(len(lines) * PDF_BODY_LEADING + 3)
            for line in lines:
                draw_text(line, x=PDF_MARGIN_X, size=PDF_BODY_SIZE, font="F2")
                y -= PDF_BODY_LEADING
            y -= 1
            body_index += 1
            continue

        prefix = "- " if tag == "li" else ""
        indent = 11 if tag == "li" else 0
        lines = _wrap_text(text, max_width=max_width - indent, font_size=PDF_BODY_SIZE)
        ensure_space(len(lines) * PDF_BODY_LEADING + 3)
        for index, line in enumerate(lines):
            draw_text(f"{prefix if index == 0 else '  '}{line}", x=PDF_MARGIN_X + indent, size=PDF_BODY_SIZE, font="F1")
            y -= PDF_BODY_LEADING
        y -= 1
        body_index += 1
    return commands


def _simple_pdf_bytes(text: str) -> bytes:
    pages: list[list[str]] = [[]]
    for command in _render_pdf_commands(text):
        if command == "__PAGE_BREAK__":
            pages.append([])
        else:
            pages[-1].append(command)

    streams = ["\n".join(page).encode("latin-1", errors="backslashreplace") for page in pages]
    page_object_ids = [5 + index * 2 for index in range(len(streams))]
    content_object_ids = [6 + index * 2 for index in range(len(streams))]
    kids = " ".join(f"{object_id} 0 R" for object_id in page_object_ids)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(streams)} >>".encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
    ]
    for page_id, content_id, stream in zip(page_object_ids, content_object_ids, streams):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {PDF_PAGE_WIDTH} {PDF_PAGE_HEIGHT}] "
            f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content_id} 0 R >>".encode("ascii")
        )
        objects.append(b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream")

    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body.extend(f"{index} 0 obj\n".encode("ascii"))
        body.extend(obj)
        body.extend(b"\nendobj\n")
    xref = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    body.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        body.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    body.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(body)


def _safe_document_html(html_content: str, *, title: str) -> str:
    return "\n".join(
        [
            "<!doctype html>",
            '<html lang="fr">',
            "<head>",
            '<meta charset="utf-8">',
            '<meta name="viewport" content="width=device-width, initial-scale=1">',
            f"<title>{html.escape(title)}</title>",
            "<style>",
            DOCUMENT_EXPORT_CSS,
            "</style>",
            "</head>",
            "<body>",
            '<main class="document-workspace">',
            '<article class="document-page">',
            f'<div class="document-body">{html_content}</div>',
            "</article>",
            "</main>",
            "</body>",
            "</html>",
        ]
    )


def _candidate_pack_files(pack_dir: Path, kind: str) -> list[Path]:
    if kind == "cv":
        prefixes = ("CV",)
    else:
        prefixes = ("LM", "Lettre", "Lettre_Motivation")
    files: list[Path] = []
    for path in sorted(pack_dir.iterdir()):
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".pdf", ".docx", ".html"}:
            continue
        if any(path.name.startswith(prefix) for prefix in prefixes):
            files.append(path)
    return files


def _first_pack_file(pack_dir: Path, kind: str, suffix: str) -> Path | None:
    suffix = suffix.lower()
    for path in _candidate_pack_files(pack_dir, kind):
        if path.suffix.lower() == suffix:
            return path
    return None


def _is_usable_pack_pdf(path: Path | None) -> bool:
    return bool(path and path.exists() and path.stat().st_size > 20_000)


def _write_pack_file(archive: zipfile.ZipFile, path: Path, *, preferred_name: str) -> None:
    archive.write(path, arcname=preferred_name)


def _replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _write_edited_docx(source_path: Path, edited_html: str, output_path: Path) -> None:
    document = Document(source_path)
    edited_blocks = extract_editable_blocks(edited_html).get("blocks", [])
    edited_texts = [str(block.get("text") or "") for block in edited_blocks]
    if not edited_texts:
        document.save(output_path)
        return

    index = 0
    for paragraph in document.paragraphs:
        if index >= len(edited_texts):
            break
        if not paragraph.text.strip():
            continue
        _replace_paragraph_text(paragraph, edited_texts[index])
        index += 1

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if index >= len(edited_texts):
                        break
                    if not paragraph.text.strip():
                        continue
                    _replace_paragraph_text(paragraph, edited_texts[index])
                    index += 1

    document.save(output_path)


def _write_docx_for_export(
    archive: zipfile.ZipFile,
    source_docx: Path,
    *,
    preferred_name: str,
    edited_html: str,
    is_dirty: bool,
    output_dir: Path,
) -> None:
    if not is_dirty:
        _write_pack_file(archive, source_docx, preferred_name=preferred_name)
        return

    patched_path = output_dir / preferred_name
    _write_edited_docx(source_docx, edited_html, patched_path)
    archive.write(patched_path, arcname=preferred_name)


def export_final_zip(
    preview_id: str,
    preview_root: str | Path,
    output_root: str | Path,
    *,
    cv_edited: str = "",
    cv_is_dirty: bool = False,
    lm_edited: str = "",
    lm_is_dirty: bool = False,
) -> Path:
    session = load_preview_session(preview_id, preview_root)
    final_cv = cv_edited if cv_is_dirty else session["cv_generated"]
    final_lm = lm_edited if lm_is_dirty else session["lm_generated"]
    output = Path(output_root)
    output.mkdir(parents=True, exist_ok=True)
    zip_path = output / "ResumeForge_documents_finaux.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        pack_dir = Path(session.get("pack_dir") or "")
        cv_pdf = _first_pack_file(pack_dir, "cv", ".pdf") if pack_dir.is_dir() and not cv_is_dirty else None
        lm_pdf = _first_pack_file(pack_dir, "lm", ".pdf") if pack_dir.is_dir() and not lm_is_dirty else None
        cv_docx = _first_pack_file(pack_dir, "cv", ".docx") if pack_dir.is_dir() else None
        lm_docx = _first_pack_file(pack_dir, "lm", ".docx") if pack_dir.is_dir() else None

        if cv_docx:
            _write_docx_for_export(
                archive,
                cv_docx,
                preferred_name="CV_Lucas_Pertusa.docx",
                edited_html=final_cv,
                is_dirty=cv_is_dirty,
                output_dir=output,
            )
        elif _is_usable_pack_pdf(cv_pdf):
            _write_pack_file(archive, cv_pdf, preferred_name="CV_Lucas_Pertusa.pdf")
        elif cv_is_dirty:
            archive.writestr("CV_Lucas_Pertusa.pdf", _simple_pdf_bytes(final_cv))
        else:
            archive.writestr("CV_Lucas_Pertusa.html", _safe_document_html(final_cv, title="CV_Lucas_Pertusa"))

        if lm_docx:
            _write_docx_for_export(
                archive,
                lm_docx,
                preferred_name="Lettre_Motivation_Lucas_Pertusa.docx",
                edited_html=final_lm,
                is_dirty=lm_is_dirty,
                output_dir=output,
            )
        elif _is_usable_pack_pdf(lm_pdf):
            _write_pack_file(archive, lm_pdf, preferred_name="Lettre_Motivation_Lucas_Pertusa.pdf")
        elif lm_is_dirty:
            archive.writestr("Lettre_Motivation_Lucas_Pertusa.pdf", _simple_pdf_bytes(final_lm))
        else:
            archive.writestr(
                "Lettre_Motivation_Lucas_Pertusa.html",
                _safe_document_html(final_lm, title="Lettre_Motivation_Lucas_Pertusa"),
            )
    return zip_path
