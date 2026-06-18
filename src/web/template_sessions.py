from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
import shutil
import subprocess
import uuid
import zipfile

from docx import Document
from docx.shared import Pt


DOCX_FILENAMES = {
    "cv": "CV_Lucas_Pertusa.docx",
    "lm": "Lettre_Motivation_Lucas_Pertusa.docx",
}
PDF_FILENAMES = {
    "cv": "CV_Lucas_Pertusa.pdf",
    "lm": "Lettre_Motivation_Lucas_Pertusa.pdf",
}


def _first_file(pack_dir: Path, pattern: str) -> Path | None:
    matches = sorted(path for path in pack_dir.glob(pattern) if path.is_file())
    return matches[0] if matches else None


def _paragraph_texts(path: Path | None) -> list[str]:
    if not path or not path.exists():
        return []
    document = Document(path)
    texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                texts.append(" | ".join(cells))
    return texts


def cv_data_from_docx(path: Path | None) -> dict:
    texts = _paragraph_texts(path)
    name = texts[0] if len(texts) > 0 else "Lucas Pertusa"
    title = texts[1] if len(texts) > 1 else ""
    contact = texts[2] if len(texts) > 2 else ""
    profile = texts[3] if len(texts) > 3 else ""
    lower_texts = [text.lower() for text in texts]
    skills_start = next((index for index, text in enumerate(lower_texts) if "compétence" in text or "skill" in text), len(texts))
    experience_start = next((index for index, text in enumerate(lower_texts) if "expérience" in text or "experience" in text), 4)
    education_start = next((index for index, text in enumerate(lower_texts) if "formation" in text or "education" in text), len(texts))

    experience_items = texts[experience_start + 1 : min(skills_start, education_start)]
    experiences = []
    if experience_items:
        company = experience_items[0] if len(experience_items) > 0 else ""
        role = experience_items[1] if len(experience_items) > 1 else ""
        dates = experience_items[2] if len(experience_items) > 2 else ""
        bullets = experience_items[3:] or experience_items[:]
        experiences.append({"company": company, "role": role, "dates": dates, "bullets": bullets[:6]})

    skills = texts[skills_start + 1 : education_start] if skills_start < len(texts) else []
    if len(skills) == 1 and "," in skills[0]:
        skills = [item.strip() for item in skills[0].split(",") if item.strip()]
    education = texts[education_start + 1 :] if education_start < len(texts) else []
    return {
        "name": name,
        "title": title,
        "contact": contact,
        "profile": profile,
        "skills": skills[:12],
        "experiences": experiences,
        "education": education[:6],
    }


def lm_data_from_docx(path: Path | None) -> dict:
    texts = _paragraph_texts(path)
    return {
        "recipient": texts[1] if len(texts) > 1 else "Madame, Monsieur,",
        "subject": texts[0] if texts else "",
        "intro": texts[2] if len(texts) > 2 else "",
        "body_1": texts[3] if len(texts) > 3 else "",
        "body_2": texts[4] if len(texts) > 4 else "",
        "closing": texts[5] if len(texts) > 5 else "",
        "signature": texts[6] if len(texts) > 6 else "Lucas Pertusa",
    }


def _set_default_font(document: Document) -> None:
    styles = document.styles
    if "Normal" in styles:
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10)


def render_cv_docx(data: dict, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _set_default_font(document)
    document.add_heading(str(data.get("name") or "Lucas Pertusa"), level=0)
    if data.get("title"):
        document.add_paragraph(str(data["title"]))
    if data.get("contact"):
        document.add_paragraph(str(data["contact"]))
    if data.get("profile"):
        document.add_heading("Profil", level=1)
        document.add_paragraph(str(data["profile"]))
    if data.get("experiences"):
        document.add_heading("Expérience", level=1)
        for experience in data.get("experiences", []):
            heading = " · ".join(
                item
                for item in [
                    str(experience.get("company") or "").strip(),
                    str(experience.get("role") or "").strip(),
                    str(experience.get("dates") or "").strip(),
                ]
                if item
            )
            if heading:
                document.add_paragraph(heading)
            for bullet in experience.get("bullets", []):
                if str(bullet).strip():
                    document.add_paragraph(str(bullet).strip(), style="List Bullet")
    if data.get("skills"):
        document.add_heading("Compétences", level=1)
        document.add_paragraph(", ".join(str(skill).strip() for skill in data.get("skills", []) if str(skill).strip()))
    if data.get("education"):
        document.add_heading("Formation", level=1)
        for item in data.get("education", []):
            if str(item).strip():
                document.add_paragraph(str(item).strip())
    document.save(output)
    return output


def render_lm_docx(data: dict, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _set_default_font(document)
    for key in ["subject", "recipient", "intro", "body_1", "body_2", "closing", "signature"]:
        value = str(data.get(key) or "").strip()
        if value:
            paragraph = document.add_paragraph(value)
            if key == "subject":
                for run in paragraph.runs:
                    run.bold = True
    document.save(output)
    return output


def _libreoffice_binary() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def convert_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path | None:
    binary = _libreoffice_binary()
    if not binary:
        return None
    output_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [binary, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        timeout=60,
    )
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    return pdf_path if pdf_path.exists() else None


def _write_json(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _split_lines(value: str) -> list[str]:
    return [line.strip() for line in str(value or "").splitlines() if line.strip()]


def _render_session_files(root: Path, cv_data: dict, lm_data: dict) -> tuple[dict, list[str]]:
    warnings: list[str] = []
    cv_docx = render_cv_docx(cv_data, root / DOCX_FILENAMES["cv"])
    lm_docx = render_lm_docx(lm_data, root / DOCX_FILENAMES["lm"])
    documents = {
        "cv": {"label": "CV", "docx_filename": DOCX_FILENAMES["cv"], "pdf_filename": ""},
        "lm": {"label": "Lettre de motivation", "docx_filename": DOCX_FILENAMES["lm"], "pdf_filename": ""},
    }
    for kind, docx_path in {"cv": cv_docx, "lm": lm_docx}.items():
        try:
            pdf_path = convert_docx_to_pdf(docx_path, root)
        except (subprocess.SubprocessError, OSError) as exc:
            pdf_path = None
            warnings.append(f"Conversion PDF {documents[kind]['label']} impossible : {exc}")
        if pdf_path:
            target = root / PDF_FILENAMES[kind]
            if pdf_path != target:
                pdf_path.replace(target)
            documents[kind]["pdf_filename"] = PDF_FILENAMES[kind]
    if not any(document["pdf_filename"] for document in documents.values()):
        warnings.append("LibreOffice indisponible : preview PDF non générée. Le téléchargement DOCX reste disponible.")
    return documents, warnings


def _session_payload(session_id: str, root: Path, documents: dict, warnings: list[str]) -> dict:
    return {
        "session_id": session_id,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "cv_data": _read_json(root / "cv_data.json"),
        "lm_data": _read_json(root / "lm_data.json"),
        "documents": documents,
        "warnings": warnings,
    }


def create_template_session(pack_dir: str | Path, session_root: str | Path) -> dict:
    pack = Path(pack_dir)
    if not pack.is_dir():
        raise ValueError("Pack de candidature introuvable.")
    cv_data = cv_data_from_docx(_first_file(pack, "CV*.docx"))
    lm_data = lm_data_from_docx(_first_file(pack, "LM*.docx"))
    return create_template_session_from_data(cv_data, lm_data, session_root)


def create_template_session_from_data(cv_data: dict, lm_data: dict, session_root: str | Path) -> dict:
    session_id = uuid.uuid4().hex
    root = Path(session_root) / session_id
    root.mkdir(parents=True, exist_ok=True)
    _write_json(root / "cv_data.json", cv_data)
    _write_json(root / "lm_data.json", lm_data)
    documents, warnings = _render_session_files(root, cv_data, lm_data)
    payload = _session_payload(session_id, root, documents, warnings)
    _write_json(root / "session.json", payload)
    return payload


def load_template_session(session_id: str, session_root: str | Path) -> dict:
    root = Path(session_root) / str(session_id)
    session_path = root / "session.json"
    if not session_path.exists():
        raise ValueError("Session template introuvable.")
    payload = _read_json(session_path)
    payload["cv_data"] = _read_json(root / "cv_data.json")
    payload["lm_data"] = _read_json(root / "lm_data.json")
    return payload


def update_template_session(session_id: str, session_root: str | Path, updates: dict) -> dict:
    root = Path(session_root) / str(session_id)
    if not root.exists():
        raise ValueError("Session template introuvable.")
    cv_data = _read_json(root / "cv_data.json")
    lm_data = _read_json(root / "lm_data.json")

    cv_updates = updates.get("cv", {})
    for key in ["name", "title", "contact", "profile"]:
        if key in cv_updates:
            cv_data[key] = cv_updates[key]
    if "skills_text" in cv_updates:
        cv_data["skills"] = _split_lines(cv_updates["skills_text"])
    if "education_text" in cv_updates:
        cv_data["education"] = _split_lines(cv_updates["education_text"])
    for index, experience in enumerate(cv_data.get("experiences", [])):
        prefix = f"experience_{index}_"
        for field in ["company", "role", "dates"]:
            form_key = f"{prefix}{field}"
            if form_key in cv_updates:
                experience[field] = cv_updates[form_key]
        bullets_key = f"{prefix}bullets"
        if bullets_key in cv_updates:
            experience["bullets"] = _split_lines(cv_updates[bullets_key])

    lm_updates = updates.get("lm", {})
    for key in ["recipient", "subject", "intro", "body_1", "body_2", "closing", "signature"]:
        if key in lm_updates:
            lm_data[key] = lm_updates[key]

    _write_json(root / "cv_data.json", cv_data)
    _write_json(root / "lm_data.json", lm_data)
    documents, warnings = _render_session_files(root, cv_data, lm_data)
    payload = _session_payload(session_id, root, documents, warnings)
    _write_json(root / "session.json", payload)
    return payload


def export_template_session_zip(session_id: str, session_root: str | Path, output_root: str | Path) -> Path:
    root = Path(session_root) / str(session_id)
    if not root.exists():
        raise ValueError("Session template introuvable.")
    output = Path(output_root)
    output.mkdir(parents=True, exist_ok=True)
    zip_path = output / "ResumeForge_documents_finaux.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for filename in [DOCX_FILENAMES["cv"], DOCX_FILENAMES["lm"]]:
            path = root / filename
            if path.exists():
                archive.write(path, arcname=filename)
    return zip_path
